"""Report export helpers (CSV / Excel) for the Data Comparison UI."""

from __future__ import annotations

import csv
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from seatunnel_agent.agents.data.text2sql.exporter import safe_stem

from .comparator import CompareReport

# ---------------------------------------------------------------------------
# Export helpers — read from CompareReport
# ---------------------------------------------------------------------------

def _download_dir() -> Path:
    """Staging dir for gr.DownloadButton files.

    Must live under the system temp dir (or the cwd): Gradio 6 refuses to
    serve anything else (InvalidPathError), which silently broke the old
    Desktop-based export — the click looked fine and no file ever arrived
    (caught when automating checklist items Q1/Q2). The browser download
    then puts the file wherever the user chose, so Desktop is not needed."""
    p = Path(tempfile.gettempdir()) / "seatunnel_agent_exports"
    p.mkdir(parents=True, exist_ok=True)
    return p


def _export_report_csv(report: CompareReport) -> str:
    """Export a CompareReport to a multi-section CSV file."""
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    fname = f"{safe_stem('data_comparison')}_{ts}.csv"
    target = _download_dir() / fname

    with open(target, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)

        if report.schema is not None:
            w.writerow(["[Schema Diff]"])
            w.writerow(["Column", "Type A", "Type B", "Status"])
            for item in report.schema.items:
                w.writerow([item.column, item.type_a, item.type_b, item.status])
            if not report.schema.items:
                w.writerow(["(identical)"])
            w.writerow([])

        if report.row_count is not None:
            w.writerow(["[Row Count]"])
            w.writerow(["Table A", "Count A", "Table B", "Count B", "Delta", "Delta %"])
            r = report.row_count
            w.writerow([r.table_a, r.count_a, r.table_b, r.count_b, r.delta, f"{r.delta_pct:.1f}%"])
            w.writerow([])

        if report.batch_counts:
            w.writerow(["[Batch Row Counts]"])
            w.writerow(["Table", "Count A", "Count B", "Delta", "Delta %"])
            for r in report.batch_counts:
                w.writerow([r.table_a, r.count_a, r.count_b, r.delta, f"{r.delta_pct:.1f}%"])
            w.writerow([])

        if report.sample is not None and report.sample.has_changes:
            w.writerow(["[Sample Data Diff]"])
            w.writerow(["Added Rows", "Removed Rows"])
            w.writerow([len(report.sample.added_rows), len(report.sample.removed_rows)])
            w.writerow([])

        if report.aggregate is not None:
            w.writerow(["[Aggregates]"])
            w.writerow(["Column", "Metric", "Value A", "Value B", "Match"])
            for item in report.aggregate.items:
                w.writerow([item.column, item.metric, item.value_a, item.value_b, item.match])

        if report.profile is not None:
            w.writerow([])
            w.writerow(["[Column Profile]"])
            w.writerow(["Column", "Distinct A", "Distinct B", "Null% A", "Null% B",
                         "Min A", "Min B", "Max A", "Max B"])
            for item in report.profile.items:
                w.writerow([item.column, item.distinct_a, item.distinct_b,
                           f"{item.null_rate_a:.1f}%", f"{item.null_rate_b:.1f}%",
                           item.min_a, item.min_b, item.max_a, item.max_b])

        if report.skew is not None and report.skew.items:
            w.writerow([])
            w.writerow(["[Skew Analysis]"])
            w.writerow(["Column", "Gini A", "Gini B", "CV A", "CV B", "Top1% A", "Top1% B"])
            for item in report.skew.items:
                w.writerow([item.column, item.gini_a, item.gini_b,
                           item.cv_a, item.cv_b, item.top1_pct_a, item.top1_pct_b])

        if report.keyed_diff is not None:
            w.writerow([])
            w.writerow(["[Keyed Diff]"])
            w.writerow(["Added Rows", "Removed Rows", "Modified Rows"])
            w.writerow([len(report.keyed_diff.added), len(report.keyed_diff.removed),
                        len(report.keyed_diff.modified)])

        if report.checksum is not None:
            w.writerow([])
            w.writerow(["[Checksum]"])
            w.writerow(["Segment", "Checksum A", "Checksum B", "Match"])
            for item in report.checksum.items:
                w.writerow([item.segment, item.checksum_a, item.checksum_b, item.match])

        if report.partition is not None and report.partition.items:
            w.writerow([])
            w.writerow(["[Partition]"])
            w.writerow(["Partition", "Count A", "Count B", "Delta"])
            for item in report.partition.items:
                w.writerow([item.partition_value, item.count_a, item.count_b, item.delta])

        if report.custom_agg is not None and report.custom_agg.items:
            w.writerow([])
            w.writerow(["[Custom Aggregates]"])
            w.writerow(["Alias", "Value A", "Value B", "Match"])
            for item in report.custom_agg.items:
                w.writerow([item.alias, item.value_a, item.value_b, item.match])

    return str(target.resolve())


def _export_report_excel(report: CompareReport) -> str:
    """Export a CompareReport to a multi-sheet Excel workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise RuntimeError("openpyxl is required for Excel export.")

    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    fname = f"{safe_stem('data_comparison')}_{ts}.xlsx"
    target = _download_dir() / fname

    wb = Workbook()
    bold = Font(bold=True)
    green_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    red_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")

    def _write_sheet(ws, headers, rows, color_col=None):
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = bold
            cell.alignment = Alignment(horizontal="center")
        for ri, row in enumerate(rows, 2):
            for ci, val in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=val)
            if color_col is not None and len(row) > color_col:
                fill = green_fill if row[color_col] else red_fill
                for ci in range(1, len(row) + 1):
                    ws.cell(row=ri, column=ci).fill = fill
        ws.freeze_panes = "A2"

    first = True
    if report.schema is not None:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Schema Diff"
        headers = ["Column", "Type A", "Type B", "Status"]
        rows = [[i.column, i.type_a, i.type_b, i.status] for i in report.schema.items]
        if not rows:
            rows = [["(identical)", "", "", ""]]
        _write_sheet(ws, headers, rows)

    if report.row_count is not None:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Row Count"
        r = report.row_count
        _write_sheet(ws, ["Table A", "Count A", "Table B", "Count B", "Delta", "Delta %"],
                     [[r.table_a, r.count_a, r.table_b, r.count_b, r.delta, f"{r.delta_pct:.1f}%"]])

    if report.batch_counts:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Batch Counts"
        rows = [[r.table_a, r.count_a, r.count_b, r.delta, f"{r.delta_pct:.1f}%", r.delta == 0]
                for r in report.batch_counts]
        _write_sheet(ws, ["Table", "Count A", "Count B", "Delta", "Delta %", "Match"], rows, color_col=5)

    if report.sample is not None and report.sample.has_changes:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Sample Diff"
        _write_sheet(ws, ["Added Rows", "Removed Rows"],
                     [[len(report.sample.added_rows), len(report.sample.removed_rows)]])

    if report.aggregate is not None and report.aggregate.items:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Aggregates"
        rows = [[i.column, i.metric, i.value_a, i.value_b, i.match] for i in report.aggregate.items]
        _write_sheet(ws, ["Column", "Metric", "Value A", "Value B", "Match"], rows, color_col=4)

    if report.profile is not None and report.profile.items:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Column Profile"
        rows = [[i.column, i.distinct_a, i.distinct_b,
                 f"{i.null_rate_a:.1f}%", f"{i.null_rate_b:.1f}%",
                 str(i.min_a) if i.min_a is not None else "",
                 str(i.min_b) if i.min_b is not None else "",
                 str(i.max_a) if i.max_a is not None else "",
                 str(i.max_b) if i.max_b is not None else ""]
                for i in report.profile.items]
        _write_sheet(ws, ["Column", "Distinct A", "Distinct B", "Null% A", "Null% B",
                          "Min A", "Min B", "Max A", "Max B"], rows)

    if report.skew is not None and report.skew.items:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Skew Analysis"
        rows = [[i.column, i.gini_a, i.gini_b, i.cv_a, i.cv_b, i.top1_pct_a, i.top1_pct_b]
                for i in report.skew.items]
        _write_sheet(ws, ["Column", "Gini A", "Gini B", "CV A", "CV B", "Top1% A", "Top1% B"], rows)

    if report.keyed_diff is not None:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Keyed Diff"
        _write_sheet(ws, ["Added Rows", "Removed Rows", "Modified Rows"],
                     [[len(report.keyed_diff.added), len(report.keyed_diff.removed),
                       len(report.keyed_diff.modified)]])

    if report.checksum is not None:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Checksum"
        rows = [[i.segment, i.checksum_a, i.checksum_b, i.match] for i in report.checksum.items]
        _write_sheet(ws, ["Segment", "Checksum A", "Checksum B", "Match"], rows, color_col=3)

    if report.partition is not None and report.partition.items:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Partition"
        rows = [[i.partition_value, i.count_a, i.count_b, i.delta] for i in report.partition.items]
        _write_sheet(ws, ["Partition", "Count A", "Count B", "Delta"], rows)

    if report.custom_agg is not None and report.custom_agg.items:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Custom Aggregates"
        rows = [[i.alias, i.value_a, i.value_b, i.match] for i in report.custom_agg.items]
        _write_sheet(ws, ["Alias", "Value A", "Value B", "Match"], rows, color_col=3)

    if first:
        wb.active.title = "Empty"
        wb.active.cell(row=1, column=1, value="No comparison data")

    wb.save(target)
    return str(target.resolve())
