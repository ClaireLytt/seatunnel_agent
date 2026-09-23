"""Gradio page for the Data Comparison agent.

Rendered inside the multipage app built by ``ui.create_ui`` via
``app.route("Data Comparison", "/datacompare")``.  No LLM — pure
deterministic comparison of two data sources.
"""

from __future__ import annotations

import csv
import json
import logging
import re
import threading
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import gradio as gr

from .text2sql.executor import (
    DS_DEFAULTS,
    DS_TYPES,
    DIALECT_NAMES,
    ENV_PREFIX,
    DatabaseConfig,
    config_from_env,
    create_executor,
)
from .text2sql.differ import diff_results, ResultDiff
from .text2sql.exporter import default_desktop_dir, safe_stem
from .data_comparison.comparator import (
    GINI_SKEW_THRESHOLD,
    AggregateResult,
    BatchFullItem,
    ChecksumResult,
    CompareReport,
    CustomAggResult,
    KeyedDiffResult,
    PartitionResult,
    ProfileItem,
    ProfileResult,
    QualityResult,
    RowCountResult,
    SchemaDiffResult,
    SkewItem,
    SkewResult,
    ThresholdConfig,
    apply_column_mapping,
    build_aggregate_sql,
    build_checksum_sql,
    build_count_sql,
    build_custom_agg_sql,
    build_incremental_sql,
    build_partition_count_sql,
    build_profile_sql,
    build_random_sample_sql,
    build_skew_sql,
    build_stratified_sample_sql,
    check_trend_alerts,
    parse_alert_rules,
    TrendAlert,
    build_row_count_result,
    build_sample_sql,
    build_trend_data,
    check_aggregate_threshold,
    check_quality_rules,
    check_row_count_threshold,
    compare_aggregates,
    compare_checksums,
    compare_custom_aggs,
    compare_partitions,
    compare_profiles,
    compare_schemas,
    compare_skew,
    detect_sensitive_columns,
    diff_by_key,
    find_common_tables,
    generate_diff_sql,
    generate_sync_config,
    is_numeric_type,
    mask_rows,
    parse_column_mapping,
    parse_custom_agg_expressions,
    parse_quality_rules,
    parse_threshold,
    run_parallel,
    ReportDiff,
    ReportDiffItem,
    diff_reports,
    build_expression_check_sql,
    get_upstream_tables,
)
from .data_comparison.i18n import dc
from .data_comparison.presets import ConnectionPresetsStore
from .data_comparison.templates_store import ComparisonTemplatesStore


# ---------------------------------------------------------------------------
# Matplotlib import (lazy — only used for trend charts)
# ---------------------------------------------------------------------------

def _get_matplotlib():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    return plt


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _esc_html(s: str) -> str:
    return (str(s).replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;"))


def _error_html(lang: str, exc: Exception) -> str:
    return f'<div style="color:#dc2626;">❌ {dc(lang, "dc_error")}: {_esc_html(str(exc))}</div>'


_log = logging.getLogger(__name__)

# Tunable limits (previously scattered magic numbers)
_SAMPLE_LIMIT = 100          # rows sampled per side
_STRATIFIED_PER_GROUP = 10   # rows per stratum in stratified sampling
_SKEW_MAX_COLUMNS = 10       # columns analysed for skew
_SKEW_MAX_ROWS = 50          # top-N frequency rows fetched per column
_MAX_AGG_COLUMNS = 20        # numeric columns per aggregate comparison
_MAX_REPORT_FILES = 50       # saved reports listed in dropdowns
_SUMMARY_MAX_ITEMS = 10      # items shown in report summary sections
_DELTA_WARN_PCT = 10         # row-count delta % below which it's a warning
_DEFAULT_SCHEDULE_MIN = 15   # default scheduled-comparison interval (minutes)

_DS_CHOICES = [DIALECT_NAMES[d] for d in DS_TYPES]
_DS_LABEL_TO_KEY = {v: k for k, v in DIALECT_NAMES.items()}
_NEEDS_AUTH = frozenset({"mysql", "sqlserver", "sparksql", "clickhouse", "doris", "postgresql"})
_NEEDS_HOST = frozenset(DS_TYPES) - {"flinksql", "sqlite"}

_DDL_KEYWORDS = re.compile(
    r"\b(DROP|DELETE|INSERT|UPDATE|ALTER|CREATE|TRUNCATE|EXEC|EXECUTE|GRANT|REVOKE"
    r"|MERGE|CALL|UNION|SET|LOAD|COPY|RENAME)\b"
    r"|INTO\s+OUTFILE",
    re.IGNORECASE,
)

_REPORTS_DIR = Path.home() / ".seatunnel-agent" / "reports"
_PRESETS_STORE = ConnectionPresetsStore()
_TEMPLATES_STORE = ComparisonTemplatesStore()
_SAMPLE_STRATEGIES = ["TOP N", "RANDOM", "STRATIFIED"]

_TH = 'style="text-align:left;padding:4px 6px;font-size:11px;border-bottom:1px solid #e5e7eb;"'
_TD = 'style="padding:4px 6px;font-size:11px;border-bottom:1px solid #f3f4f6;"'


def _details_card(summary_html: str, body_html: str, border_color: str = "#e0e7ff") -> str:
    return (
        f'<details open style="border:1px solid {border_color};border-radius:8px;'
        f'padding:10px;background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;cursor:pointer;">'
        f'{summary_html}</summary>{body_html}</details>'
    )


def _build_webhook_summary(report: CompareReport) -> dict:
    """Extract top-10 schema diffs and agg mismatches as structured data."""
    summary: dict[str, Any] = {}
    if report.schema and report.schema.has_changes:
        summary["schema_diffs"] = [
            {"column": it.column, "status": it.status,
             "type_a": it.type_a, "type_b": it.type_b}
            for it in report.schema.items[:_SUMMARY_MAX_ITEMS]
        ]
    if report.aggregate and report.aggregate.mismatches > 0:
        summary["agg_mismatches"] = [
            {"column": it.column, "metric": it.metric,
             "value_a": it.value_a, "value_b": it.value_b}
            for it in report.aggregate.items[:_SUMMARY_MAX_ITEMS]
            if not it.match
        ][:_SUMMARY_MAX_ITEMS]
    return summary


def _send_webhook(url: str, payload: dict) -> tuple[bool, str]:
    """POST JSON payload to a webhook URL. Returns (success, message)."""
    import urllib.request
    import urllib.error
    data = json.dumps(payload, ensure_ascii=False, default=str).encode("utf-8")
    req = urllib.request.Request(
        url, data=data,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            return True, f"{resp.status}"
    except (urllib.error.URLError, urllib.error.HTTPError, OSError) as exc:
        return False, str(exc)


def _validate_where(clause: str, lang: str) -> tuple[bool, str]:
    """Validate a WHERE clause — reject DDL keywords."""
    if not clause.strip():
        return True, ""
    m = _DDL_KEYWORDS.search(clause)
    if m:
        return False, dc(lang, "dc_where_invalid").format(kw=m.group(0))
    return True, ""


def _with_validation(holder, holder_lock, fn):
    """Wrap a compare callback with connection/table checks and error handling."""
    def wrapper(*args):
        lang_val = args[2] if len(args) > 2 else "en"
        with holder_lock:
            if holder.get("executor_a") is None or holder.get("executor_b") is None:
                return dc(lang_val, "dc_connect_both")
        table_a = args[0] if len(args) > 0 else ""
        table_b = args[1] if len(args) > 1 else ""
        if not table_a or not table_b:
            return dc(lang_val, "dc_select_tables")
        try:
            return fn(*args)
        except Exception as e:
            return _error_html(lang_val, e)
    return wrapper


# ---------------------------------------------------------------------------
# HTML card builders
# ---------------------------------------------------------------------------

def build_schema_diff_card(result: SchemaDiffResult, lang: str = "en") -> str:
    t = lambda k: dc(lang, k)
    esc = _esc_html

    if not result.has_changes:
        return (
            '<div style="border:1px solid #bbf7d0;border-radius:8px;padding:12px;'
            'background:#f0fdf4;margin-bottom:8px;">'
            f'<b style="color:#16a34a;">{t("dc_schema_result")}</b> — '
            f'{t("dc_identical")}</div>'
        )

    th = 'style="padding:4px 10px;text-align:left;border-bottom:1px solid #e5e7eb;font-size:12px;"'
    td = 'style="padding:4px 10px;font-size:12px;border-bottom:1px solid #f3f4f6;"'
    rows_html = ""
    for item in result.items:
        if item.status == "removed":
            color = "#fef2f2"
            badge = f'<span style="color:#dc2626;font-weight:600;">{t("dc_removed")}</span>'
        elif item.status == "added":
            color = "#f0fdf4"
            badge = f'<span style="color:#16a34a;font-weight:600;">{t("dc_added")}</span>'
        else:
            color = "#fffbeb"
            badge = f'<span style="color:#d97706;font-weight:600;">{t("dc_type_changed")}</span>'
        rows_html += (
            f'<tr style="background:{color};">'
            f'<td {td}>{esc(item.column)}</td>'
            f'<td {td}>{esc(item.type_a)}</td>'
            f'<td {td}>{esc(item.type_b)}</td>'
            f'<td {td}>{badge}</td></tr>'
        )

    summary = []
    if result.cols_only_a:
        summary.append(f'<span style="color:#dc2626;">{t("dc_removed")}: {result.cols_only_a}</span>')
    if result.cols_only_b:
        summary.append(f'<span style="color:#16a34a;">{t("dc_added")}: {result.cols_only_b}</span>')
    if result.type_changes:
        summary.append(f'<span style="color:#d97706;">{t("dc_type_changed")}: {result.type_changes}</span>')

    return (
        '<details open style="border:1px solid #e0e7ff;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#4338ca;cursor:pointer;">'
        f'{t("dc_schema_result")} — {" · ".join(summary)}</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_column")}</th><th {th}>{t("dc_type_a")}</th>'
        f'<th {th}>{t("dc_type_b")}</th><th {th}>{t("dc_status")}</th></tr>'
        f'{rows_html}</table></details>'
    )


def _threshold_badge(passed: bool, lang: str) -> str:
    if passed:
        return (' <span style="background:#16a34a;color:#fff;padding:2px 8px;'
                f'border-radius:4px;font-size:11px;margin-left:8px;">'
                f'{dc(lang, "dc_threshold_pass")}</span>')
    return (' <span style="background:#dc2626;color:#fff;padding:2px 8px;'
            f'border-radius:4px;font-size:11px;margin-left:8px;">'
            f'{dc(lang, "dc_threshold_fail")}</span>')


def build_count_card(
    result: RowCountResult, lang: str = "en",
    threshold: ThresholdConfig | None = None,
) -> str:
    t = lambda k: dc(lang, k)

    if result.delta == 0:
        border = "#bbf7d0"
        bg = "#f0fdf4"
        badge = f'<span style="color:#16a34a;font-weight:600;">{t("dc_match")}</span>'
    elif abs(result.delta_pct) < _DELTA_WARN_PCT:
        border = "#fde68a"
        bg = "#fffbeb"
        badge = f'<span style="color:#d97706;font-weight:600;">{t("dc_delta")}: {result.delta:+d} ({result.delta_pct:+.1f}%)</span>'
    else:
        border = "#fecaca"
        bg = "#fef2f2"
        badge = f'<span style="color:#dc2626;font-weight:600;">{t("dc_delta")}: {result.delta:+d} ({result.delta_pct:+.1f}%)</span>'

    threshold_html = ""
    if threshold is not None:
        passed = check_row_count_threshold(result, threshold)
        threshold_html = _threshold_badge(passed, lang)

    return (
        f'<div style="border:1px solid {border};border-radius:8px;padding:12px;'
        f'background:{bg};margin-bottom:8px;">'
        f'<b style="color:#0369a1;">{t("dc_count_result")}</b>{threshold_html}<br>'
        f'{t("dc_count_a")}: <b>{result.count_a:,}</b> &nbsp; '
        f'{t("dc_count_b")}: <b>{result.count_b:,}</b> &nbsp; '
        f'{badge}</div>'
    )


def build_sample_diff_card(
    diff: ResultDiff, lang: str = "en", columns: list[str] | None = None,
) -> str:
    """Build HTML card for sample data diff, including actual row preview."""
    t = lambda k: dc(lang, k)
    esc = _esc_html

    if not diff.has_changes:
        return (
            '<div style="border:1px solid #bbf7d0;border-radius:8px;padding:12px;'
            'background:#f0fdf4;margin-bottom:8px;">'
            f'<b style="color:#0369a1;">{t("dc_sample_result")}</b> — '
            f'{t("dc_no_diff")}</div>'
        )

    parts: list[str] = []
    if diff.added_rows:
        parts.append(f'<span style="color:#16a34a;">{t("dc_added_rows")}: +{len(diff.added_rows)}</span>')
    if diff.removed_rows:
        parts.append(f'<span style="color:#dc2626;">{t("dc_removed_rows")}: -{len(diff.removed_rows)}</span>')
    if diff.columns_added:
        parts.append(f'<span style="color:#16a34a;">+cols: {", ".join(esc(c) for c in diff.columns_added[:5])}</span>')
    if diff.columns_removed:
        parts.append(f'<span style="color:#dc2626;">-cols: {", ".join(esc(c) for c in diff.columns_removed[:5])}</span>')

    header = (
        '<div style="border:1px solid #bae6fd;border-radius:8px;padding:12px;'
        'background:#f0f9ff;margin-bottom:8px;">'
        f'<b style="color:#0369a1;">{t("dc_sample_result")}</b><br>'
        f'{" · ".join(parts)}'
    )

    row_tables = ""
    th = 'style="padding:3px 8px;text-align:left;border-bottom:1px solid #e5e7eb;font-size:11px;font-weight:600;"'
    td = 'style="padding:3px 8px;font-size:11px;border-bottom:1px solid #f3f4f6;"'

    preview_limit = 10
    cols = columns or []

    if diff.added_rows and cols:
        show = diff.added_rows[:preview_limit]
        row_tables += (
            f'<details style="margin-top:8px;"><summary style="font-size:12px;color:#16a34a;cursor:pointer;">'
            f'{t("dc_sample_preview")} — {t("dc_added_rows")} '
            f'({t("dc_showing_n_rows").format(n=len(show))})</summary>'
            f'<table style="width:100%;border-collapse:collapse;margin-top:4px;">'
            f'<tr>{"".join(f"<th {th}>{esc(c)}</th>" for c in cols)}</tr>'
        )
        for row in show:
            cells = "".join(
                f'<td {td}>{esc(str(row[i]) if i < len(row) else "")}</td>'
                for i in range(len(cols))
            )
            row_tables += f'<tr style="background:#f0fdf4;">{cells}</tr>'
        row_tables += "</table></details>"

    if diff.removed_rows and cols:
        show = diff.removed_rows[:preview_limit]
        row_tables += (
            f'<details style="margin-top:8px;"><summary style="font-size:12px;color:#dc2626;cursor:pointer;">'
            f'{t("dc_sample_preview")} — {t("dc_removed_rows")} '
            f'({t("dc_showing_n_rows").format(n=len(show))})</summary>'
            f'<table style="width:100%;border-collapse:collapse;margin-top:4px;">'
            f'<tr>{"".join(f"<th {th}>{esc(c)}</th>" for c in cols)}</tr>'
        )
        for row in show:
            cells = "".join(
                f'<td {td}>{esc(str(row[i]) if i < len(row) else "")}</td>'
                for i in range(len(cols))
            )
            row_tables += f'<tr style="background:#fef2f2;">{cells}</tr>'
        row_tables += "</table></details>"

    return header + row_tables + "</div>"


def build_aggregate_card(result: AggregateResult, lang: str = "en") -> str:
    t = lambda k: dc(lang, k)

    if not result.items:
        return (
            '<div style="border:1px solid #e5e7eb;border-radius:8px;padding:12px;'
            'background:#f9fafb;margin-bottom:8px;">'
            f'<b style="color:#6b7280;">{t("dc_agg_result")}</b> — '
            f'{t("dc_no_numeric")}</div>'
        )

    th = 'style="padding:4px 10px;text-align:left;border-bottom:1px solid #e5e7eb;font-size:12px;"'
    td = 'style="padding:4px 10px;font-size:12px;border-bottom:1px solid #f3f4f6;"'
    rows_html = ""
    for item in result.items:
        bg = "#fff" if item.match else "#fef2f2"
        badge = "" if item.match else f' <span style="color:#dc2626;font-size:10px;">({t("dc_mismatch")})</span>'
        va = _esc_html("NULL" if item.value_a is None else str(item.value_a))
        vb = _esc_html("NULL" if item.value_b is None else str(item.value_b))
        rows_html += (
            f'<tr style="background:{bg};">'
            f'<td {td}>{_esc_html(item.column)}</td>'
            f'<td {td}>{_esc_html(item.metric)}</td>'
            f'<td {td}>{va}</td><td {td}>{vb}</td>'
            f'<td {td}>{badge}</td></tr>'
        )

    summary_badge = (
        f'<span style="color:#dc2626;">{result.mismatches} {t("dc_mismatch")}</span>'
        if result.mismatches else
        f'<span style="color:#16a34a;">{t("dc_match")}</span>'
    )

    return (
        '<details open style="border:1px solid #e0e7ff;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#7c3aed;cursor:pointer;">'
        f'{t("dc_agg_result")} — {summary_badge}</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_column")}</th><th {th}>{t("dc_metric")}</th>'
        f'<th {th}>{t("dc_value_a")}</th><th {th}>{t("dc_value_b")}</th>'
        f'<th {th}></th></tr>'
        f'{rows_html}</table></details>'
    )


def build_batch_count_card(results: list[RowCountResult], lang: str = "en") -> str:
    """Build HTML card showing row counts for multiple table pairs."""
    t = lambda k: dc(lang, k)
    esc = _esc_html

    if not results:
        return (
            '<div style="border:1px solid #e5e7eb;border-radius:8px;padding:12px;'
            'background:#f9fafb;margin-bottom:8px;">'
            f'<b style="color:#6b7280;">{t("dc_batch_result")}</b> — '
            f'{t("dc_no_common")}</div>'
        )

    mismatches = sum(1 for r in results if r.delta != 0)
    summary = t("dc_common_tables").format(n=len(results))
    if mismatches:
        summary += f' · <span style="color:#dc2626;">{t("dc_diffs_found").format(n=mismatches)}</span>'
    else:
        summary += f' · <span style="color:#16a34a;">{t("dc_no_diffs")}</span>'

    th = 'style="padding:4px 10px;text-align:left;border-bottom:1px solid #e5e7eb;font-size:12px;"'
    td = 'style="padding:4px 10px;font-size:12px;border-bottom:1px solid #f3f4f6;"'
    rows_html = ""
    for r in results:
        if r.delta == 0:
            bg = "#f0fdf4"
            badge = f'<span style="color:#16a34a;">{t("dc_match")}</span>'
        else:
            bg = "#fef2f2"
            badge = f'<span style="color:#dc2626;">{r.delta:+d} ({r.delta_pct:+.1f}%)</span>'
        rows_html += (
            f'<tr style="background:{bg};">'
            f'<td {td}>{esc(r.table_a)}</td>'
            f'<td {td}>{r.count_a:,}</td>'
            f'<td {td}>{r.count_b:,}</td>'
            f'<td {td}>{badge}</td></tr>'
        )

    return (
        '<details open style="border:1px solid #e0e7ff;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#0369a1;cursor:pointer;">'
        f'{t("dc_batch_result")} — {summary}</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_table_name")}</th><th {th}>{t("dc_count_a")}</th>'
        f'<th {th}>{t("dc_count_b")}</th><th {th}>{t("dc_delta")}</th></tr>'
        f'{rows_html}</table></details>'
    )


def build_summary_card(report: CompareReport, lang: str = "en") -> str:
    """Build a summary card showing comparison overview."""
    t = lambda k: dc(lang, k)

    count = 0
    diffs = 0
    if report.schema is not None:
        count += 1
        if report.schema.has_changes:
            diffs += 1
    if report.row_count is not None:
        count += 1
        if report.row_count.delta != 0:
            diffs += 1
    if report.sample is not None:
        count += 1
        if report.sample.has_changes:
            diffs += 1
    if report.aggregate is not None:
        count += 1
        if report.aggregate.mismatches > 0:
            diffs += 1

    if diffs:
        border, bg = "#fde68a", "#fffbeb"
        diff_text = f'<span style="color:#d97706;font-weight:600;">{t("dc_diffs_found").format(n=diffs)}</span>'
    else:
        border, bg = "#bbf7d0", "#f0fdf4"
        diff_text = f'<span style="color:#16a34a;font-weight:600;">{t("dc_no_diffs")}</span>'

    elapsed = t("dc_elapsed").format(ms=report.elapsed_ms) if report.elapsed_ms else ""

    return (
        f'<div style="border:1px solid {border};border-radius:8px;padding:12px;'
        f'background:{bg};margin-bottom:10px;">'
        f'<b style="color:#1e40af;">{t("dc_summary")}</b> &nbsp; '
        f'{t("dc_comparisons_done").format(n=count)} · {diff_text}'
        f'{f" · {elapsed}" if elapsed else ""}</div>'
    )


# ---------------------------------------------------------------------------
# New card builders  (A — keyed diff, C — profile)
# ---------------------------------------------------------------------------

def build_keyed_diff_card(result: KeyedDiffResult, lang: str = "en") -> str:
    """Build HTML card for key-based row diff."""
    t = lambda k: dc(lang, k)
    esc = _esc_html

    if not result.added and not result.removed and not result.modified:
        return (
            '<div style="border:1px solid #bbf7d0;border-radius:8px;padding:12px;'
            'background:#f0fdf4;margin-bottom:8px;">'
            f'<b style="color:#0369a1;">{t("dc_key_diff")}</b> — '
            f'{t("dc_no_diff")}</div>'
        )

    parts: list[str] = []
    if result.added:
        parts.append(f'<span style="color:#16a34a;">{t("dc_added_rows")}: +{len(result.added)}</span>')
    if result.removed:
        parts.append(f'<span style="color:#dc2626;">{t("dc_removed_rows")}: -{len(result.removed)}</span>')
    if result.modified:
        parts.append(f'<span style="color:#d97706;">{t("dc_modified_rows")}: {len(result.modified)}</span>')

    html = (
        '<div style="border:1px solid #e0e7ff;border-radius:8px;padding:12px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<b style="color:#0369a1;">{t("dc_key_diff")}</b> '
        f'(key: {esc(", ".join(result.key_columns))})<br>'
        f'{" · ".join(parts)}'
    )

    th = 'style="padding:3px 8px;text-align:left;border-bottom:1px solid #e5e7eb;font-size:11px;font-weight:600;"'
    td = 'style="padding:3px 8px;font-size:11px;border-bottom:1px solid #f3f4f6;"'

    if result.modified:
        html += (
            f'<details style="margin-top:8px;" open><summary style="font-size:12px;color:#d97706;cursor:pointer;">'
            f'{t("dc_modified_rows")} ({len(result.modified)})</summary>'
            f'<table style="width:100%;border-collapse:collapse;margin-top:4px;">'
            f'<tr><th {th}>Key</th><th {th}>{t("dc_column")}</th>'
            f'<th {th}>{t("dc_old_val")}</th><th {th}>{t("dc_new_val")}</th></tr>'
        )
        for mr in result.modified[:50]:
            key_str = esc(str(mr.key))
            for col, old_v, new_v in mr.changes:
                html += (
                    f'<tr style="background:#fffbeb;">'
                    f'<td {td}>{key_str}</td><td {td}>{esc(col)}</td>'
                    f'<td style="padding:3px 8px;font-size:11px;border-bottom:1px solid #f3f4f6;color:#dc2626;">{esc(str(old_v))}</td>'
                    f'<td style="padding:3px 8px;font-size:11px;border-bottom:1px solid #f3f4f6;color:#16a34a;">{esc(str(new_v))}</td></tr>'
                )
        html += "</table></details>"

    # Feature F — drill-down per modified row
    if result.modified:
        for mr in result.modified[:50]:
            if mr.row_a is not None or mr.row_b is not None:
                key_str = _esc_html(str(mr.key))
                detail_th = 'style="padding:2px 6px;text-align:left;border-bottom:1px solid #e5e7eb;font-size:10px;font-weight:600;"'
                detail_td = 'style="padding:2px 6px;font-size:10px;border-bottom:1px solid #f3f4f6;"'
                cols = result.key_columns if result.key_columns else []
                row_a_vals = mr.row_a if mr.row_a else ()
                row_b_vals = mr.row_b if mr.row_b else ()
                html += (
                    f'<details style="margin-left:16px;margin-top:2px;">'
                    f'<summary style="font-size:11px;color:#6b7280;cursor:pointer;">'
                    f'{dc(lang, "dc_drill_down")} — Key: {key_str}</summary>'
                    f'<div style="display:flex;gap:12px;margin-top:4px;">'
                )
                if row_a_vals:
                    html += (
                        f'<div><b style="font-size:10px;">{dc(lang, "dc_row_values_a")}</b>'
                        f'<table style="border-collapse:collapse;">'
                    )
                    for i, v in enumerate(row_a_vals):
                        html += f'<tr><td {detail_td}>{esc(str(v))}</td></tr>'
                    html += '</table></div>'
                if row_b_vals:
                    html += (
                        f'<div><b style="font-size:10px;">{dc(lang, "dc_row_values_b")}</b>'
                        f'<table style="border-collapse:collapse;">'
                    )
                    for i, v in enumerate(row_b_vals):
                        html += f'<tr><td {detail_td}>{esc(str(v))}</td></tr>'
                    html += '</table></div>'
                html += '</div></details>'

    html += "</div>"
    return html


def build_quality_card(results: list[QualityResult], lang: str = "en") -> str:
    """Build HTML card for quality rule check results."""
    t = lambda k: dc(lang, k)
    esc = _esc_html

    if not results:
        return ""

    passed = sum(1 for r in results if r.passed)
    failed = len(results) - passed

    th = 'style="padding:4px 10px;text-align:left;border-bottom:1px solid #e5e7eb;font-size:12px;"'
    td = 'style="padding:4px 10px;font-size:12px;border-bottom:1px solid #f3f4f6;"'
    rows_html = ""
    for r in results:
        bg = "#f0fdf4" if r.passed else "#fef2f2"
        badge = (f'<span style="color:#16a34a;">{t("dc_quality_pass")}</span>'
                 if r.passed else
                 f'<span style="color:#dc2626;">{t("dc_quality_fail")}</span>')
        rows_html += (
            f'<tr style="background:{bg};">'
            f'<td {td}>{esc(r.rule.column)}</td>'
            f'<td {td}>{esc(r.rule.rule_type)}</td>'
            f'<td {td}>{r.actual_value:.2f}</td>'
            f'<td {td}>{badge}</td></tr>'
        )

    summary = (f'<span style="color:#16a34a;">{passed} {t("dc_quality_pass")}</span>'
               f' / <span style="color:#dc2626;">{failed} {t("dc_quality_fail")}</span>')

    return (
        '<details open style="border:1px solid #e0e7ff;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#0891b2;cursor:pointer;">'
        f'{t("dc_quality_result")} — {summary}</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_column")}</th><th {th}>Rule</th>'
        f'<th {th}>Actual</th><th {th}>{t("dc_status")}</th></tr>'
        f'{rows_html}</table></details>'
    )


def build_diff_sql_card(sql: str, lang: str = "en") -> str:
    """Build HTML card showing generated diff SQL."""
    t = lambda k: dc(lang, k)
    if not sql.strip():
        return (
            '<div style="border:1px solid #bbf7d0;border-radius:8px;padding:12px;'
            'background:#f0fdf4;margin-bottom:8px;">'
            f'<b style="color:#0369a1;">{t("dc_diff_sql_result")}</b> — '
            f'{t("dc_diff_sql_no_diff")}</div>'
        )
    return (
        f'<div style="border:1px solid #e0e7ff;border-radius:8px;padding:12px;'
        f'background:#f8fafc;margin-bottom:8px;">'
        f'<b style="color:#0369a1;">{t("dc_diff_sql_result")}</b>'
        f'<pre style="background:#1e293b;color:#e2e8f0;padding:16px;border-radius:8px;'
        f'overflow-x:auto;font-size:12px;margin-top:8px;white-space:pre-wrap;">'
        f'{_esc_html(sql)}</pre></div>'
    )


def build_profile_card(result: ProfileResult, lang: str = "en") -> str:
    """Build HTML card for column profile comparison."""
    t = lambda k: dc(lang, k)
    esc = _esc_html

    if not result.items:
        return (
            '<div style="border:1px solid #e5e7eb;border-radius:8px;padding:12px;'
            'background:#f9fafb;margin-bottom:8px;">'
            f'<b style="color:#6b7280;">{t("dc_profile_result")}</b></div>'
        )

    th = 'style="padding:4px 8px;text-align:left;border-bottom:1px solid #e5e7eb;font-size:11px;font-weight:600;"'
    td = 'style="padding:4px 8px;font-size:11px;border-bottom:1px solid #f3f4f6;"'

    rows_html = ""
    for item in result.items:
        dist_match = item.distinct_a == item.distinct_b
        null_match = abs(item.null_rate_a - item.null_rate_b) < 0.1
        bg = "#fff" if (dist_match and null_match) else "#fffbeb"
        rows_html += (
            f'<tr style="background:{bg};">'
            f'<td {td}>{esc(item.column)}</td>'
            f'<td {td}>{item.distinct_a:,}</td><td {td}>{item.distinct_b:,}</td>'
            f'<td {td}>{item.null_rate_a:.1f}%</td><td {td}>{item.null_rate_b:.1f}%</td>'
            f'<td {td}>{esc(str(item.min_a) if item.min_a is not None else "")}</td>'
            f'<td {td}>{esc(str(item.min_b) if item.min_b is not None else "")}</td>'
            f'<td {td}>{esc(str(item.max_a) if item.max_a is not None else "")}</td>'
            f'<td {td}>{esc(str(item.max_b) if item.max_b is not None else "")}</td></tr>'
        )

    return (
        '<details open style="border:1px solid #e0e7ff;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#0891b2;cursor:pointer;">'
        f'{t("dc_profile_result")} — {result.total_a:,} / {result.total_b:,} rows</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_column")}</th>'
        f'<th {th}>{t("dc_distinct")} A</th><th {th}>{t("dc_distinct")} B</th>'
        f'<th {th}>{t("dc_null_rate")} A</th><th {th}>{t("dc_null_rate")} B</th>'
        f'<th {th}>{t("dc_min")} A</th><th {th}>{t("dc_min")} B</th>'
        f'<th {th}>{t("dc_max")} A</th><th {th}>{t("dc_max")} B</th></tr>'
        f'{rows_html}</table></details>'
    )


def build_skew_card(result: SkewResult, lang: str = "en") -> str:
    """Build HTML card for data skew analysis."""
    t = lambda k: dc(lang, k)
    esc = _esc_html
    if not result.items:
        return (
            '<div style="padding:10px;color:#6b7280;background:#f9fafb;'
            f'border-radius:8px;margin-bottom:8px;">{t("dc_skew_result")}</div>'
        )
    th = _TH
    td = _TD

    items_html = ""
    for item in result.items:
        warn_a = ' style="color:#dc2626;font-weight:600;"' if item.gini_a > GINI_SKEW_THRESHOLD else ""
        warn_b = ' style="color:#dc2626;font-weight:600;"' if item.gini_b > GINI_SKEW_THRESHOLD else ""
        items_html += (
            f'<tr style="background:#f0fdfa;">'
            f'<td {td}><b>{esc(item.column)}</b></td>'
            f'<td {td}{warn_a}>{item.gini_a}</td><td {td}{warn_b}>{item.gini_b}</td>'
            f'<td {td}>{item.cv_a}</td><td {td}>{item.cv_b}</td>'
            f'<td {td}>{item.top1_pct_a}%</td><td {td}>{item.top1_pct_b}%</td>'
            f'<td {td}>{item.ndv_a}</td><td {td}>{item.ndv_b}</td></tr>'
        )
        for bkt in item.buckets:
            items_html += (
                f'<tr><td {td} style="padding-left:20px;font-size:10px;">'
                f'{esc(bkt.value)}</td>'
                f'<td {td} colspan="2" style="font-size:10px;">'
                f'{bkt.count_a:,} ({bkt.pct_a}%)</td>'
                f'<td {td} colspan="2" style="font-size:10px;">'
                f'{bkt.count_b:,} ({bkt.pct_b}%)</td>'
                f'<td {td} colspan="4"></td></tr>'
            )

    return (
        '<details open style="border:1px solid #e0e7ff;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#0d9488;cursor:pointer;">'
        f'{t("dc_skew_result")} — {result.total_a:,} / {result.total_b:,} rows</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_skew_column")}</th>'
        f'<th {th}>{t("dc_skew_gini")} A</th><th {th}>{t("dc_skew_gini")} B</th>'
        f'<th {th}>{t("dc_skew_cv")} A</th><th {th}>{t("dc_skew_cv")} B</th>'
        f'<th {th}>{t("dc_skew_top1")} A</th><th {th}>{t("dc_skew_top1")} B</th>'
        f'<th {th}>{t("dc_skew_ndv")} A</th><th {th}>{t("dc_skew_ndv")} B</th></tr>'
        f'{items_html}</table></details>'
    )


def build_checksum_card(result, lang: str = "en") -> str:
    """Build HTML card for checksum comparison."""
    t = lambda k: dc(lang, k)
    if not result.items:
        return (
            '<div style="padding:10px;color:#6b7280;background:#f9fafb;'
            f'border-radius:8px;margin-bottom:8px;">{t("dc_checksum_result")}</div>'
        )
    if result.mismatch_count == 0:
        badge = f'<span style="color:#16a34a;font-weight:600;">✅ {t("dc_checksum_all_match")}</span>'
    else:
        badge = (f'<span style="color:#dc2626;font-weight:600;">'
                 f'{result.mismatch_count} {t("dc_checksum_mismatch")}</span>')

    th = _TH
    td = _TD
    rows_html = ""
    for item in result.items:
        color = "#f0fdf4" if item.match else "#fef2f2"
        status = t("dc_checksum_match") if item.match else t("dc_checksum_mismatch")
        rows_html += (
            f'<tr style="background:{color};">'
            f'<td {td}>{_esc_html(str(item.segment))}</td>'
            f'<td {td}>{_esc_html(item.checksum_a[:16])}</td>'
            f'<td {td}>{_esc_html(item.checksum_b[:16])}</td>'
            f'<td {td}>{status}</td></tr>'
        )
    table_html = (
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_checksum_segment")}</th>'
        f'<th {th}>Checksum A</th><th {th}>Checksum B</th>'
        f'<th {th}>Status</th></tr>'
        f'{rows_html}</table>'
    )
    return _details_card(
        f'<span style="color:#4f46e5;">{t("dc_checksum_result")} — {badge}</span>',
        table_html,
    )


def build_partition_card(result, lang: str = "en") -> str:
    """Build HTML card for partition-level comparison."""
    t = lambda k: dc(lang, k)
    if not result.items:
        return (
            '<div style="padding:10px;color:#6b7280;background:#f9fafb;'
            f'border-radius:8px;margin-bottom:8px;">{t("dc_partition_result")}</div>'
        )
    th = _TH
    td = _TD
    rows_html = ""
    for item in result.items:
        color = "#f0fdf4" if item.delta == 0 else "#fef2f2"
        rows_html += (
            f'<tr style="background:{color};">'
            f'<td {td}>{_esc_html(item.partition_value)}</td>'
            f'<td {td}>{item.count_a:,}</td>'
            f'<td {td}>{item.count_b:,}</td>'
            f'<td {td}>{item.delta:+,}</td>'
            f'<td {td}>{item.delta_pct:+.1f}%</td></tr>'
        )
    return (
        '<details open style="border:1px solid #e0e7ff;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#7c3aed;cursor:pointer;">'
        f'{t("dc_partition_result")} — {t("dc_partition_total_delta")}: {result.total_delta:,}</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_partition_value")}</th>'
        f'<th {th}>Count A</th><th {th}>Count B</th>'
        f'<th {th}>{t("dc_partition_delta")}</th><th {th}>%</th></tr>'
        f'{rows_html}</table></details>'
    )


def build_custom_agg_card(result, lang: str = "en") -> str:
    """Build HTML card for custom aggregate comparison."""
    t = lambda k: dc(lang, k)
    if not result.items:
        return (
            '<div style="padding:10px;color:#6b7280;background:#f9fafb;'
            f'border-radius:8px;margin-bottom:8px;">{t("dc_custom_agg_result")}</div>'
        )
    th = _TH
    td = _TD
    rows_html = ""
    for item in result.items:
        color = "#f0fdf4" if item.match else "#fef2f2"
        rows_html += (
            f'<tr style="background:{color};">'
            f'<td {td}>{_esc_html(item.alias)}</td>'
            f'<td {td}>{_esc_html(item.expression)}</td>'
            f'<td {td}>{_esc_html(str(item.value_a))}</td>'
            f'<td {td}>{_esc_html(str(item.value_b))}</td>'
            f'<td {td}>{"✓" if item.match else "✗"}</td></tr>'
        )
    mismatch_badge = ""
    if result.mismatches > 0:
        mismatch_badge = f' — <span style="color:#dc2626;">{result.mismatches} mismatch</span>'
    table_html = (
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_custom_agg_alias")}</th>'
        f'<th {th}>{t("dc_custom_agg_expression")}</th>'
        f'<th {th}>A</th><th {th}>B</th><th {th}>Match</th></tr>'
        f'{rows_html}</table>'
    )
    return _details_card(
        f'<span style="color:#ea580c;">{t("dc_custom_agg_result")}{mismatch_badge}</span>',
        table_html,
    )


def build_batch_template_card(results: list[dict], lang: str = "en") -> str:
    """Build HTML card showing batch template run results."""
    t = lambda k: dc(lang, k)
    if not results:
        return ""
    th = _TH
    td = _TD
    rows = ""
    pass_n = sum(1 for r in results if r["status"] == "pass")
    fail_n = len(results) - pass_n
    for r in results:
        color = "#f0fdf4" if r["status"] == "pass" else "#fef2f2"
        status_text = t("dc_batch_templates_pass") if r["status"] == "pass" else t("dc_batch_templates_fail")
        detail = r.get("detail", "")
        cnt_info = ""
        if "count_a" in r:
            cnt_info = f'{r["count_a"]} / {r["count_b"]}'
        sc = r.get("schema_changes", "")
        am = r.get("agg_mismatches", "")
        rows += (
            f'<tr style="background:{color};">'
            f'<td {td}>{_esc_html(r["name"])}</td>'
            f'<td {td}>{_esc_html(r.get("table_a", ""))}</td>'
            f'<td {td}>{_esc_html(r.get("table_b", ""))}</td>'
            f'<td {td}>{cnt_info}</td>'
            f'<td {td}>{sc}</td>'
            f'<td {td}>{am}</td>'
            f'<td {td}><b>{status_text}</b></td>'
            f'<td {td}>{_esc_html(detail)}</td></tr>'
        )
    summary = dc(lang, "dc_batch_summary").format(total=len(results), pass_n=pass_n, fail_n=fail_n)
    return (
        '<details open style="border:1px solid #c7d2fe;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#6366f1;cursor:pointer;">'
        f'{t("dc_batch_templates_result")} — {summary}</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>Template</th><th {th}>Table A</th><th {th}>Table B</th>'
        f'<th {th}>Counts</th><th {th}>{t("dc_batch_tpl_schema")}</th>'
        f'<th {th}>{t("dc_batch_tpl_agg")}</th>'
        f'<th {th}>Status</th><th {th}>Detail</th></tr>'
        f'{rows}</table></details>'
    )


def build_trend_alert_card(alerts: list[TrendAlert], lang: str = "en") -> str:
    """Build HTML card showing trend alert results."""
    t = lambda k: dc(lang, k)
    if not alerts:
        return ""
    th = _TH
    td = _TD
    rows = ""
    for a in alerts:
        color = "#fef2f2" if a.triggered else "#f0fdf4"
        status = t("dc_alert_triggered") if a.triggered else t("dc_alert_ok")
        rows += (
            f'<tr style="background:{color};">'
            f'<td {td}>{_esc_html(a.metric)}</td>'
            f'<td {td}>{a.threshold}</td>'
            f'<td {td}>{a.consecutive_actual}/{a.consecutive_required}</td>'
            f'<td {td}>{a.current_value}</td>'
            f'<td {td}><b>{status}</b></td></tr>'
        )
    return (
        '<details open style="border:1px solid #fde68a;border-radius:8px;padding:10px;'
        'background:#fffbeb;margin-bottom:8px;margin-top:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#d97706;cursor:pointer;">'
        f'{t("dc_alert_result")}</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_alert_metric")}</th>'
        f'<th {th}>{t("dc_alert_threshold")}</th>'
        f'<th {th}>{t("dc_alert_consecutive")}</th>'
        f'<th {th}>Value</th><th {th}>Status</th></tr>'
        f'{rows}</table></details>'
    )


def build_report_diff_card(rdiff: ReportDiff, lang: str = "en") -> str:
    t = lambda k: dc(lang, k)
    if not rdiff.has_changes:
        return (
            f'<div style="border:1px solid #86efac;border-radius:8px;padding:12px;'
            f'background:#f0fdf4;margin-bottom:8px;">'
            f'<b style="color:#16a34a;">{t("dc_report_diff_no_change")}</b></div>'
        )
    th = _TH
    td = _TD
    rows = ""
    for item in rdiff.items:
        if not item.changed:
            continue
        rows += (
            f'<tr><td {td}>{_esc_html(item.section)}</td>'
            f'<td {td}>{_esc_html(item.field)}</td>'
            f'<td {td}>{item.old_value}</td>'
            f'<td {td}>{item.new_value}</td></tr>'
        )
    return (
        '<details open style="border:1px solid #c7d2fe;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#6366f1;cursor:pointer;">'
        f'{t("dc_report_diff_result")}</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_report_diff_section")}</th>'
        f'<th {th}>{t("dc_report_diff_field")}</th>'
        f'<th {th}>{t("dc_report_diff_old_val")}</th>'
        f'<th {th}>{t("dc_report_diff_new_val")}</th></tr>'
        f'{rows}</table></details>'
    )


def build_lineage_card(table: str, upstream: list[str], lang: str = "en") -> str:
    t = lambda k: dc(lang, k)
    if not upstream:
        return (
            f'<div style="border:1px solid #e5e7eb;border-radius:8px;padding:12px;'
            f'background:#f9fafb;margin-bottom:8px;">'
            f'<b>{t("dc_lineage_title")}</b>: {t("dc_lineage_no_deps")}</div>'
        )
    items_html = "".join(f"<li>{_esc_html(t)}</li>" for t in upstream)
    return (
        '<details open style="border:1px solid #c7d2fe;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#6366f1;cursor:pointer;">'
        f'{dc(lang, "dc_lineage_title")} — {_esc_html(table)}</summary>'
        f'<ul style="margin:8px 0;padding-left:20px;">{items_html}</ul></details>'
    )


# ---------------------------------------------------------------------------
# Standalone report page  (opens in new tab)
# ---------------------------------------------------------------------------

def build_standalone_report(report: CompareReport, lang: str = "en") -> str:
    """Build a full standalone HTML page for the comparison report."""
    t = lambda k: dc(lang, k)
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    cards: list[str] = []
    cards.append(build_summary_card(report, lang))
    if report.schema is not None:
        cards.append(build_schema_diff_card(report.schema, lang))
    if report.row_count is not None:
        cards.append(build_count_card(report.row_count, lang))
    if report.sample is not None:
        cards.append(build_sample_diff_card(report.sample, lang))
    if report.keyed_diff is not None:
        cards.append(build_keyed_diff_card(report.keyed_diff, lang))
    if report.aggregate is not None:
        cards.append(build_aggregate_card(report.aggregate, lang))
    if report.profile is not None:
        cards.append(build_profile_card(report.profile, lang))
    if report.skew is not None:
        cards.append(build_skew_card(report.skew, lang))
    if report.checksum is not None:
        cards.append(build_checksum_card(report.checksum, lang))
    if report.partition is not None:
        cards.append(build_partition_card(report.partition, lang))
    if report.custom_agg is not None:
        cards.append(build_custom_agg_card(report.custom_agg, lang))
    if report.batch_counts is not None:
        cards.append(build_batch_count_card(report.batch_counts, lang))

    body = "\n".join(cards)

    return f"""<!DOCTYPE html>
<html lang="{lang}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{t("dc_report_title")}</title>
<style>
  * {{ box-sizing: border-box; }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto,
                 'Helvetica Neue', Arial, 'PingFang SC', 'Microsoft YaHei', sans-serif;
    max-width: 1100px; margin: 0 auto; padding: 24px 20px;
    background: #f8fafc; color: #1e293b; line-height: 1.5;
  }}
  .report-header {{
    display: flex; justify-content: space-between; align-items: center;
    border-bottom: 2px solid #e2e8f0; padding-bottom: 12px; margin-bottom: 20px;
  }}
  .report-header h1 {{ font-size: 22px; margin: 0; color: #1e40af; }}
  .report-header .meta {{ font-size: 12px; color: #64748b; text-align: right; }}
  .print-btn {{
    padding: 6px 16px; background: #3b82f6; color: #fff; border: none;
    border-radius: 6px; cursor: pointer; font-size: 13px; margin-left: 12px;
  }}
  .print-btn:hover {{ background: #2563eb; }}
  table {{ font-size: 12px; }}
  details {{ margin-bottom: 10px; }}
  @media print {{
    .print-btn {{ display: none; }}
    body {{ background: #fff; padding: 10px; }}
    details {{ break-inside: avoid; }}
  }}
</style>
</head>
<body>
<div class="report-header">
  <div>
    <h1>{t("dc_report_title")}</h1>
    <div class="meta">{ts}</div>
  </div>
  <button class="print-btn" onclick="window.print()">{t("dc_print")}</button>
</div>
{body}
</body>
</html>"""


# ---------------------------------------------------------------------------
# Batch full comparison card  (F)
# ---------------------------------------------------------------------------

def build_batch_full_card(
    items: list[BatchFullItem], lang: str = "en",
    threshold: ThresholdConfig | None = None,
) -> str:
    t = lambda k: dc(lang, k)
    esc = _esc_html

    passed = sum(1 for i in items if not i.has_diffs)
    failed = len(items) - passed

    summary = t("dc_batch_summary").format(
        pass_n=passed, fail_n=failed, total=len(items))

    th = 'style="padding:4px 10px;text-align:left;border-bottom:1px solid #e5e7eb;font-size:12px;"'
    td = 'style="padding:4px 10px;font-size:12px;border-bottom:1px solid #f3f4f6;"'

    rows_html = ""
    for item in items:
        bg = "#fef2f2" if item.has_diffs else "#f0fdf4"
        schema_ok = "✅" if (item.schema and not item.schema.has_changes) else "❌" if item.schema else "—"
        delta_str = f"{item.row_count.delta:+d}" if item.row_count else "—"
        agg_ok = "✅" if (item.aggregate and item.aggregate.mismatches == 0) else (
            f"❌ {item.aggregate.mismatches}" if item.aggregate else "—")
        status = f'<span style="color:#16a34a;">✅</span>' if not item.has_diffs else (
            f'<span style="color:#dc2626;">❌</span>')
        th_badge = ""
        if threshold and item.row_count:
            th_badge = _threshold_badge(check_row_count_threshold(item.row_count, threshold), lang)
        rows_html += (
            f'<tr style="background:{bg};">'
            f'<td {td}>{esc(item.table_a)}</td>'
            f'<td {td}>{schema_ok}</td>'
            f'<td {td}>{delta_str}{th_badge}</td>'
            f'<td {td}>{agg_ok}</td>'
            f'<td {td}>{status}</td></tr>'
        )

    return (
        '<details open style="border:1px solid #e0e7ff;border-radius:8px;padding:10px;'
        'background:#f8fafc;margin-bottom:8px;">'
        f'<summary style="font-weight:600;font-size:13px;color:#7c3aed;cursor:pointer;">'
        f'{t("dc_batch_full_result")} — {summary}</summary>'
        f'<table style="width:100%;border-collapse:collapse;margin-top:8px;">'
        f'<tr><th {th}>{t("dc_table_name")}</th>'
        f'<th {th}>{t("dc_compare_schema")}</th>'
        f'<th {th}>{t("dc_delta")}</th>'
        f'<th {th}>{t("dc_compare_agg")}</th>'
        f'<th {th}>{t("dc_status")}</th></tr>'
        f'{rows_html}</table></details>'
    )


# ---------------------------------------------------------------------------
# Export helpers — read from CompareReport
# ---------------------------------------------------------------------------

def _export_report_csv(report: CompareReport) -> str:
    """Export a CompareReport to a multi-section CSV file."""
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    fname = f"{safe_stem('data_comparison')}_{ts}.csv"
    target = default_desktop_dir() / fname

    with open(target, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)

        if report.schema is not None:
            w.writerow(["[Schema Diff]"])
            w.writerow(["Column", "Type A", "Type B", "Status"])
            for item in report.schema.items:
                w.writerow([item.column, item.type_a, item.type_b, item.status])
            if not report.schema.items:
                w.writerow(["(identical)"])
            w.writerow([])

        if report.row_count is not None:
            w.writerow(["[Row Count]"])
            w.writerow(["Table A", "Count A", "Table B", "Count B", "Delta", "Delta %"])
            r = report.row_count
            w.writerow([r.table_a, r.count_a, r.table_b, r.count_b, r.delta, f"{r.delta_pct:.1f}%"])
            w.writerow([])

        if report.batch_counts:
            w.writerow(["[Batch Row Counts]"])
            w.writerow(["Table", "Count A", "Count B", "Delta", "Delta %"])
            for r in report.batch_counts:
                w.writerow([r.table_a, r.count_a, r.count_b, r.delta, f"{r.delta_pct:.1f}%"])
            w.writerow([])

        if report.sample is not None and report.sample.has_changes:
            w.writerow(["[Sample Data Diff]"])
            w.writerow(["Added Rows", "Removed Rows"])
            w.writerow([len(report.sample.added_rows), len(report.sample.removed_rows)])
            w.writerow([])

        if report.aggregate is not None:
            w.writerow(["[Aggregates]"])
            w.writerow(["Column", "Metric", "Value A", "Value B", "Match"])
            for item in report.aggregate.items:
                w.writerow([item.column, item.metric, item.value_a, item.value_b, item.match])

        if report.profile is not None:
            w.writerow([])
            w.writerow(["[Column Profile]"])
            w.writerow(["Column", "Distinct A", "Distinct B", "Null% A", "Null% B",
                         "Min A", "Min B", "Max A", "Max B"])
            for item in report.profile.items:
                w.writerow([item.column, item.distinct_a, item.distinct_b,
                           f"{item.null_rate_a:.1f}%", f"{item.null_rate_b:.1f}%",
                           item.min_a, item.min_b, item.max_a, item.max_b])

        if report.skew is not None and report.skew.items:
            w.writerow([])
            w.writerow(["[Skew Analysis]"])
            w.writerow(["Column", "Gini A", "Gini B", "CV A", "CV B", "Top1% A", "Top1% B"])
            for item in report.skew.items:
                w.writerow([item.column, item.gini_a, item.gini_b,
                           item.cv_a, item.cv_b, item.top1_pct_a, item.top1_pct_b])

        if report.keyed_diff is not None:
            w.writerow([])
            w.writerow(["[Keyed Diff]"])
            w.writerow(["Added Rows", "Removed Rows", "Modified Rows"])
            w.writerow([len(report.keyed_diff.added), len(report.keyed_diff.removed),
                        len(report.keyed_diff.modified)])

        if report.checksum is not None:
            w.writerow([])
            w.writerow(["[Checksum]"])
            w.writerow(["Segment", "Checksum A", "Checksum B", "Match"])
            for item in report.checksum.items:
                w.writerow([item.segment, item.checksum_a, item.checksum_b, item.match])

        if report.partition is not None and report.partition.items:
            w.writerow([])
            w.writerow(["[Partition]"])
            w.writerow(["Partition", "Count A", "Count B", "Delta"])
            for item in report.partition.items:
                w.writerow([item.partition_value, item.count_a, item.count_b, item.delta])

        if report.custom_agg is not None and report.custom_agg.items:
            w.writerow([])
            w.writerow(["[Custom Aggregates]"])
            w.writerow(["Alias", "Value A", "Value B", "Match"])
            for item in report.custom_agg.items:
                w.writerow([item.alias, item.value_a, item.value_b, item.match])

    return str(target.resolve())


def _export_report_excel(report: CompareReport) -> str:
    """Export a CompareReport to a multi-sheet Excel workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        raise RuntimeError("openpyxl is required for Excel export.")

    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    fname = f"{safe_stem('data_comparison')}_{ts}.xlsx"
    target = default_desktop_dir() / fname

    wb = Workbook()
    bold = Font(bold=True)
    green_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    red_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")

    def _write_sheet(ws, headers, rows, color_col=None):
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = bold
            cell.alignment = Alignment(horizontal="center")
        for ri, row in enumerate(rows, 2):
            for ci, val in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=val)
            if color_col is not None and len(row) > color_col:
                fill = green_fill if row[color_col] else red_fill
                for ci in range(1, len(row) + 1):
                    ws.cell(row=ri, column=ci).fill = fill
        ws.freeze_panes = "A2"

    first = True
    if report.schema is not None:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Schema Diff"
        headers = ["Column", "Type A", "Type B", "Status"]
        rows = [[i.column, i.type_a, i.type_b, i.status] for i in report.schema.items]
        if not rows:
            rows = [["(identical)", "", "", ""]]
        _write_sheet(ws, headers, rows)

    if report.row_count is not None:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Row Count"
        r = report.row_count
        _write_sheet(ws, ["Table A", "Count A", "Table B", "Count B", "Delta", "Delta %"],
                     [[r.table_a, r.count_a, r.table_b, r.count_b, r.delta, f"{r.delta_pct:.1f}%"]])

    if report.batch_counts:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Batch Counts"
        rows = [[r.table_a, r.count_a, r.count_b, r.delta, f"{r.delta_pct:.1f}%", r.delta == 0]
                for r in report.batch_counts]
        _write_sheet(ws, ["Table", "Count A", "Count B", "Delta", "Delta %", "Match"], rows, color_col=5)

    if report.sample is not None and report.sample.has_changes:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Sample Diff"
        _write_sheet(ws, ["Added Rows", "Removed Rows"],
                     [[len(report.sample.added_rows), len(report.sample.removed_rows)]])

    if report.aggregate is not None and report.aggregate.items:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Aggregates"
        rows = [[i.column, i.metric, i.value_a, i.value_b, i.match] for i in report.aggregate.items]
        _write_sheet(ws, ["Column", "Metric", "Value A", "Value B", "Match"], rows, color_col=4)

    if report.profile is not None and report.profile.items:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Column Profile"
        rows = [[i.column, i.distinct_a, i.distinct_b,
                 f"{i.null_rate_a:.1f}%", f"{i.null_rate_b:.1f}%",
                 str(i.min_a) if i.min_a is not None else "",
                 str(i.min_b) if i.min_b is not None else "",
                 str(i.max_a) if i.max_a is not None else "",
                 str(i.max_b) if i.max_b is not None else ""]
                for i in report.profile.items]
        _write_sheet(ws, ["Column", "Distinct A", "Distinct B", "Null% A", "Null% B",
                          "Min A", "Min B", "Max A", "Max B"], rows)

    if report.skew is not None and report.skew.items:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Skew Analysis"
        rows = [[i.column, i.gini_a, i.gini_b, i.cv_a, i.cv_b, i.top1_pct_a, i.top1_pct_b]
                for i in report.skew.items]
        _write_sheet(ws, ["Column", "Gini A", "Gini B", "CV A", "CV B", "Top1% A", "Top1% B"], rows)

    if report.keyed_diff is not None:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Keyed Diff"
        _write_sheet(ws, ["Added Rows", "Removed Rows", "Modified Rows"],
                     [[len(report.keyed_diff.added), len(report.keyed_diff.removed),
                       len(report.keyed_diff.modified)]])

    if report.checksum is not None:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Checksum"
        rows = [[i.segment, i.checksum_a, i.checksum_b, i.match] for i in report.checksum.items]
        _write_sheet(ws, ["Segment", "Checksum A", "Checksum B", "Match"], rows, color_col=3)

    if report.partition is not None and report.partition.items:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Partition"
        rows = [[i.partition_value, i.count_a, i.count_b, i.delta] for i in report.partition.items]
        _write_sheet(ws, ["Partition", "Count A", "Count B", "Delta"], rows)

    if report.custom_agg is not None and report.custom_agg.items:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Custom Aggregates"
        rows = [[i.alias, i.value_a, i.value_b, i.match] for i in report.custom_agg.items]
        _write_sheet(ws, ["Alias", "Value A", "Value B", "Match"], rows, color_col=3)

    if first:
        wb.active.title = "Empty"
        wb.active.cell(row=1, column=1, value="No comparison data")

    wb.save(target)
    return str(target.resolve())


# ---------------------------------------------------------------------------
# Page renderer
# ---------------------------------------------------------------------------

def render_data_comparison_page(app=None) -> None:  # noqa: C901
    """Render the Data Comparison page components."""
    from dotenv import load_dotenv
    load_dotenv()

    lang = "en"
    t = lambda k: dc(lang, k)

    holder: dict[str, Any] = {
        "executor_a": None,
        "executor_b": None,
        "tables_a": [],
        "tables_b": [],
        "last_report": None,
        "schedule_timer": None,
        "schedule_active": False,
        "schedule_last_run": None,
    }
    holder_lock = threading.Lock()

    # ── Connection helper (shared by A / B) ──
    def _do_connect(ds_label, host, port, db, username, password, lang_val, side):
        t_fn = lambda k: dc(lang_val, k)

        def _err(msg):
            return msg, gr.update(), gr.update(), gr.update(), gr.update()

        ds_type = _DS_LABEL_TO_KEY.get(ds_label)
        if ds_type is None:
            return _err(f"❌ {t_fn('dc_error')}: unknown datasource {_esc_html(str(ds_label))}")
        # Hidden Gradio textboxes (auth/host fields for engines that don't
        # need them) submit None from the browser, not "".
        h = (host or "").strip()
        p = (port or "").strip()
        d = (db or "").strip()
        u = (username or "").strip() or None
        pw = (password or "").strip() or None
        defaults = DS_DEFAULTS[ds_type]
        default_port = str(defaults["port"])
        default_db = defaults["database"]

        if ds_type == "sqlite":
            cfg = DatabaseConfig(
                ds_type="sqlite", host="", port=0,
                database=d or str(default_db),
            )
        else:
            if not h:
                fallback = config_from_env(ds_type)
                if fallback:
                    h, p, d = fallback.host, str(fallback.port), fallback.database
                    u, pw = fallback.username, fallback.password

            if not h:
                var = f"{ENV_PREFIX.get(ds_type, ds_type.upper())}_HOST"
                return _err(f"❌ {t_fn('dc_host_required').format(var=var)}")

            try:
                port_int = int(p or default_port)
            except ValueError:
                return _err(f"❌ {t_fn('dc_port_not_number')}")

            cfg = DatabaseConfig(
                ds_type=ds_type,
                host=h,
                port=port_int,
                database=d or str(default_db),
                username=u,
                password=pw,
            )
        try:
            executor = create_executor(cfg)
        except Exception as e:
            return _err(f"❌ {_esc_html(str(e))}")
        ok, msg = executor.test_connection()
        if not ok:
            return _err(f"❌ {_esc_html(msg)}")

        try:
            tables = executor.show_tables()
        except Exception as e:
            return _err(f"❌ {_esc_html(str(e))}")

        with holder_lock:
            holder[f"executor_{side}"] = executor
            holder[f"tables_{side}"] = tables

        dialect = DIALECT_NAMES.get(ds_type, ds_type)
        status = f"✅ {dialect} {msg} · {t_fn('dc_tables_loaded').format(n=len(tables))}"
        # Backfill connection fields so env/default fallbacks are visible
        if ds_type == "sqlite":
            host_upd, port_upd = gr.update(), gr.update()
        else:
            host_upd, port_upd = gr.update(value=cfg.host), gr.update(value=str(cfg.port))
        return (status, gr.update(choices=tables, value=None),
                host_upd, port_upd, gr.update(value=cfg.database))

    # ── Table search filter (G) ──

    def _filter_tables(search_text, side):
        with holder_lock:
            all_tables = list(holder[f"tables_{side}"])
        if not search_text.strip():
            return gr.update(choices=all_tables)
        q = search_text.strip().lower()
        filtered = [t for t in all_tables if q in t.lower()]
        return gr.update(choices=filtered)

    # ── Inner comparison logic (no validation — used by both single + all) ──

    def _snap_executors():
        with holder_lock:
            return holder["executor_a"], holder["executor_b"]

    def _schema_inner(table_a, table_b, lang_val, where_val="", desc_a=None, desc_b=None):
        ex_a, ex_b = _snap_executors()
        if desc_a is None and desc_b is None:
            desc_a, desc_b = run_parallel(
                lambda: ex_a.describe_table(table_a),
                lambda: ex_b.describe_table(table_b),
            )
        else:
            if desc_a is None:
                desc_a = ex_a.describe_table(table_a)
            if desc_b is None:
                desc_b = ex_b.describe_table(table_b)
        cols_a = [{"name": c.name, "type": c.dtype} for c in desc_a.columns]
        cols_b = [{"name": c.name, "type": c.dtype} for c in desc_b.columns]
        result = compare_schemas(cols_a, cols_b, table_a, table_b)
        return result, build_schema_diff_card(result, lang_val)

    def _count_inner(table_a, table_b, lang_val, where_val=""):
        ex_a, ex_b = _snap_executors()
        sql_a = build_count_sql(table_a, where_val)
        sql_b = build_count_sql(table_b, where_val)
        ra, rb = run_parallel(
            lambda: ex_a.run(sql_a, max_rows=1),
            lambda: ex_b.run(sql_b, max_rows=1),
        )
        ca = int(ra.rows[0][0]) if ra.rows else 0
        cb = int(rb.rows[0][0]) if rb.rows else 0
        result = build_row_count_result(table_a, ca, table_b, cb)
        return result, build_count_card(result, lang_val)

    def _sample_inner(table_a, table_b, lang_val, where_val="", key_cols_str="",
                      strategy="TOP N", col_mapping_str="", masking_on=True,
                      stratified_col=""):
        ex_a, ex_b = _snap_executors()
        limit = _SAMPLE_LIMIT
        stratified_col = stratified_col or ""  # hidden textbox submits None
        if strategy == "STRATIFIED" and stratified_col.strip():
            sql_a = build_stratified_sample_sql(table_a, stratified_col.strip(),
                                                _STRATIFIED_PER_GROUP,
                                                ex_a.config.ds_type, where_val)
            sql_b = build_stratified_sample_sql(table_b, stratified_col.strip(),
                                                _STRATIFIED_PER_GROUP,
                                                ex_b.config.ds_type, where_val)
        elif strategy == "RANDOM":
            sql_a = build_random_sample_sql(table_a, limit, ex_a.config.ds_type, where_val)
            sql_b = build_random_sample_sql(table_b, limit, ex_b.config.ds_type, where_val)
        else:
            sql_a = build_sample_sql(table_a, limit, where_val, ds_type=ex_a.config.ds_type)
            sql_b = build_sample_sql(table_b, limit, where_val, ds_type=ex_b.config.ds_type)
        ra, rb = run_parallel(
            lambda: ex_a.run(sql_a, max_rows=limit),
            lambda: ex_b.run(sql_b, max_rows=limit),
        )

        mapping = parse_column_mapping(col_mapping_str) if col_mapping_str else {}
        if mapping:
            rb_cols = apply_column_mapping(rb.columns, mapping)
        else:
            rb_cols = rb.columns

        rows_a = list(ra.rows)
        rows_b = list(rb.rows)

        if masking_on:
            sens_a = detect_sensitive_columns(ra.columns, rows_a)
            sens_b = detect_sensitive_columns(rb_cols, rows_b)
            if sens_a:
                rows_a = mask_rows(ra.columns, rows_a, sens_a)
            if sens_b:
                rows_b = mask_rows(rb_cols, rows_b, sens_b)

        key_cols = [k.strip() for k in key_cols_str.split(",") if k.strip()] if key_cols_str else []

        fallback_notice = ""
        if key_cols:
            try:
                kd_result = diff_by_key(ra.columns, rows_a, rb_cols, rows_b, key_cols)
                return kd_result, build_keyed_diff_card(kd_result, lang_val)
            except (ValueError, IndexError) as exc:
                _log.warning("Keyed diff failed for %s / %s, falling back to "
                             "positional diff", table_a, table_b, exc_info=True)
                fallback_notice = (
                    '<div style="border:1px solid #fde68a;border-radius:8px;'
                    'padding:8px 12px;background:#fffbeb;margin-bottom:8px;'
                    'color:#d97706;font-size:12px;">⚠️ '
                    f'{dc(lang_val, "dc_keyed_diff_fallback").format(err=_esc_html(str(exc)))}'
                    '</div>'
                )

        diff = diff_results(ra.columns, rows_a, rb_cols, rows_b)
        cols = ra.columns or rb_cols
        return diff, fallback_notice + build_sample_diff_card(diff, lang_val, columns=cols)

    def _agg_inner(table_a, table_b, lang_val, where_val="", desc_a=None, desc_b=None):
        ex_a, ex_b = _snap_executors()
        if desc_a is None and desc_b is None:
            desc_a, desc_b = run_parallel(
                lambda: ex_a.describe_table(table_a),
                lambda: ex_b.describe_table(table_b),
            )
        else:
            if desc_a is None:
                desc_a = ex_a.describe_table(table_a)
            if desc_b is None:
                desc_b = ex_b.describe_table(table_b)
        names_a = {c.name.lower(): c for c in desc_a.columns}
        names_b = {c.name.lower(): c for c in desc_b.columns}
        shared_numeric = [
            c.name for key, c in names_a.items()
            if key in names_b and is_numeric_type(c.dtype)
        ]
        if not shared_numeric:
            return AggregateResult(table_a, table_b), (
                '<div style="border:1px solid #e5e7eb;border-radius:8px;padding:12px;'
                'background:#f9fafb;margin-bottom:8px;">'
                f'<b style="color:#6b7280;">{dc(lang_val, "dc_agg_result")}</b> — '
                f'{dc(lang_val, "dc_no_numeric")}</div>'
            )
        sql_a = build_aggregate_sql(table_a, shared_numeric, where_val,
                                    ds_type=ex_a.config.ds_type)
        sql_b = build_aggregate_sql(table_b, shared_numeric, where_val,
                                    ds_type=ex_b.config.ds_type)
        ra, rb = run_parallel(
            lambda: ex_a.run(sql_a, max_rows=1),
            lambda: ex_b.run(sql_b, max_rows=1),
        )
        row_a = ra.rows[0] if ra.rows else ()
        row_b = rb.rows[0] if rb.rows else ()
        result = compare_aggregates(table_a, row_a, table_b, row_b, shared_numeric)
        return result, build_aggregate_card(result, lang_val)

    def _profile_inner(table_a, table_b, lang_val, where_val=""):
        ex_a, ex_b = _snap_executors()
        desc_a, desc_b = run_parallel(
            lambda: ex_a.describe_table(table_a),
            lambda: ex_b.describe_table(table_b),
        )
        names_a = {c.name.lower(): c.name for c in desc_a.columns}
        names_b = {c.name.lower(): c.name for c in desc_b.columns}
        shared_cols = [names_a[k] for k in names_a if k in names_b]

        if not shared_cols:
            return ProfileResult(table_a, table_b), build_profile_card(
                ProfileResult(table_a, table_b), lang_val)

        sql_a = build_profile_sql(table_a, shared_cols, where_val,
                                  ds_type=ex_a.config.ds_type)
        sql_b = build_profile_sql(table_b, shared_cols, where_val,
                                  ds_type=ex_b.config.ds_type)
        ra, rb = run_parallel(
            lambda: ex_a.run(sql_a, max_rows=1),
            lambda: ex_b.run(sql_b, max_rows=1),
        )
        row_a = ra.rows[0] if ra.rows else ()
        row_b = rb.rows[0] if rb.rows else ()
        total_a = int(row_a[0]) if row_a else 0
        total_b = int(row_b[0]) if row_b else 0
        result = compare_profiles(table_a, row_a, total_a, table_b, row_b, total_b, shared_cols)
        return result, build_profile_card(result, lang_val)

    # ── Wrapped callbacks ──

    def _store_partial(key: str, value: Any) -> None:
        """Store a single comparison result into the report."""
        with holder_lock:
            if holder["last_report"] is None:
                holder["last_report"] = CompareReport()
            setattr(holder["last_report"], key, value)

    def _compare_schema_fn(table_a, table_b, lang_val, where_val):
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        result, html = _schema_inner(table_a, table_b, lang_val, where_val)
        _store_partial("schema", result)
        return html

    def _compare_count_fn(table_a, table_b, lang_val, where_val, threshold_str=""):
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        threshold = parse_threshold(threshold_str)
        result, html = _count_inner(table_a, table_b, lang_val, where_val)
        html = build_count_card(result, lang_val, threshold)
        _store_partial("row_count", result)
        return html

    def _compare_sample_fn(table_a, table_b, lang_val, where_val, key_cols_str,
                           strategy="TOP N", col_mapping_str="", masking_on=True,
                           stratified_col=""):
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        result, html = _sample_inner(
            table_a, table_b, lang_val, where_val, key_cols_str,
            strategy, col_mapping_str, masking_on, stratified_col)
        if isinstance(result, KeyedDiffResult):
            _store_partial("keyed_diff", result)
        else:
            _store_partial("sample", result)
        return html

    def _compare_agg_fn(table_a, table_b, lang_val, where_val):
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        result, html = _agg_inner(table_a, table_b, lang_val, where_val)
        _store_partial("aggregate", result)
        return html

    def _compare_profile_fn(table_a, table_b, lang_val, where_val):
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        result, html = _profile_inner(table_a, table_b, lang_val, where_val)
        _store_partial("profile", result)
        return html

    # ── Skew analysis (BB) ──

    def _skew_inner(table_a, table_b, lang_val, where_val="", skew_cols_str=""):
        ex_a, ex_b = _snap_executors()
        cols = [c.strip() for c in skew_cols_str.split(",") if c.strip()]
        if not cols:
            return SkewResult(table_a, table_b), build_skew_card(
                SkewResult(table_a, table_b), lang_val)

        cols = cols[:_SKEW_MAX_COLUMNS]
        count_sql_a = build_count_sql(table_a, where_val)
        count_sql_b = build_count_sql(table_b, where_val)
        ra_cnt, rb_cnt = run_parallel(
            lambda: ex_a.run(count_sql_a, max_rows=1),
            lambda: ex_b.run(count_sql_b, max_rows=1),
        )
        total_a = int(ra_cnt.rows[0][0]) if ra_cnt.rows else 0
        total_b = int(rb_cnt.rows[0][0]) if rb_cnt.rows else 0

        from concurrent.futures import ThreadPoolExecutor as _LocalPool

        def _process_skew_col(col):
            sql_a = build_skew_sql(table_a, col, where_val, ds_type=ex_a.config.ds_type)
            sql_b = build_skew_sql(table_b, col, where_val, ds_type=ex_b.config.ds_type)
            ra, rb = run_parallel(
                lambda s=sql_a: ex_a.run(s, max_rows=_SKEW_MAX_ROWS),
                lambda s=sql_b: ex_b.run(s, max_rows=_SKEW_MAX_ROWS),
            )
            return compare_skew(
                table_a, table_b, col,
                ra.rows, total_a,
                rb.rows, total_b,
            )

        with _LocalPool(max_workers=min(len(cols), 5)) as col_pool:
            futures = [col_pool.submit(_process_skew_col, c) for c in cols]
            items: list[SkewItem] = [f.result() for f in futures]

        result = SkewResult(
            table_a=table_a, table_b=table_b,
            total_a=total_a, total_b=total_b,
            items=items,
        )
        return result, build_skew_card(result, lang_val)

    def _compare_skew_fn(table_a, table_b, lang_val, where_val, skew_cols_str):
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        if not skew_cols_str or not skew_cols_str.strip():
            return f'<div style="color:#dc2626;">{dc(lang_val, "dc_skew_no_columns")}</div>'
        result, html = _skew_inner(table_a, table_b, lang_val, where_val, skew_cols_str)
        _store_partial("skew", result)
        return html

    _compare_schema = _with_validation(holder, holder_lock, _compare_schema_fn)
    _compare_count = _with_validation(holder, holder_lock, _compare_count_fn)
    _compare_sample = _with_validation(holder, holder_lock, _compare_sample_fn)
    _compare_agg = _with_validation(holder, holder_lock, _compare_agg_fn)
    _compare_profile = _with_validation(holder, holder_lock, _compare_profile_fn)
    _compare_skew = _with_validation(holder, holder_lock, _compare_skew_fn)

    # ── Checksum comparison (CC) ──

    def _checksum_inner(table_a, table_b, lang_val, where_val="", checksum_cols_str=""):
        ex_a, ex_b = _snap_executors()
        cols = [c.strip() for c in checksum_cols_str.split(",") if c.strip()]
        if not cols:
            desc_a, desc_b = run_parallel(
                lambda: ex_a.describe_table(table_a),
                lambda: ex_b.describe_table(table_b),
            )
            names_a = {c.name.lower() for c in desc_a.columns}
            names_b = {c.name.lower() for c in desc_b.columns}
            cols = [c.name for c in desc_a.columns if c.name.lower() in names_b][:_MAX_AGG_COLUMNS]
        if not cols:
            return ChecksumResult(table_a, table_b), build_checksum_card(
                ChecksumResult(table_a, table_b), lang_val)
        sql_a = build_checksum_sql(table_a, cols, ex_a.config.ds_type, where_val)
        sql_b = build_checksum_sql(table_b, cols, ex_b.config.ds_type, where_val)
        ra, rb = run_parallel(
            lambda: ex_a.run(sql_a, max_rows=100),
            lambda: ex_b.run(sql_b, max_rows=100),
        )
        result = compare_checksums(table_a, table_b, ra.rows, rb.rows)
        return result, build_checksum_card(result, lang_val)

    def _compare_checksum_fn(table_a, table_b, lang_val, where_val, checksum_cols_str=""):
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        result, html = _checksum_inner(table_a, table_b, lang_val, where_val, checksum_cols_str)
        _store_partial("checksum", result)
        return html

    _compare_checksum = _with_validation(holder, holder_lock, _compare_checksum_fn)

    # ── Partition comparison (DD) ──

    def _partition_inner(table_a, table_b, lang_val, where_val="", partition_col=""):
        ex_a, ex_b = _snap_executors()
        if not partition_col.strip():
            return PartitionResult(table_a, table_b), build_partition_card(
                PartitionResult(table_a, table_b), lang_val)
        sql_a = build_partition_count_sql(table_a, partition_col.strip(), where_val,
                                          ds_type=ex_a.config.ds_type)
        sql_b = build_partition_count_sql(table_b, partition_col.strip(), where_val,
                                          ds_type=ex_b.config.ds_type)
        ra, rb = run_parallel(
            lambda: ex_a.run(sql_a, max_rows=1000),
            lambda: ex_b.run(sql_b, max_rows=1000),
        )
        result = compare_partitions(table_a, table_b, partition_col.strip(),
                                    ra.rows, rb.rows)
        return result, build_partition_card(result, lang_val)

    def _compare_partition_fn(table_a, table_b, lang_val, where_val, partition_col):
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        if not partition_col or not partition_col.strip():
            return f'<div style="color:#dc2626;">{dc(lang_val, "dc_partition_col_hint")}</div>'
        result, html = _partition_inner(table_a, table_b, lang_val, where_val, partition_col)
        _store_partial("partition", result)
        return html

    _compare_partition = _with_validation(holder, holder_lock, _compare_partition_fn)

    # ── Custom aggregates (EE) ──

    def _custom_agg_inner(table_a, table_b, lang_val, where_val="", expressions_str=""):
        ex_a, ex_b = _snap_executors()
        expressions = parse_custom_agg_expressions(expressions_str)
        if not expressions:
            return CustomAggResult(table_a, table_b), build_custom_agg_card(
                CustomAggResult(table_a, table_b), lang_val)
        sql_a = build_custom_agg_sql(table_a, expressions, where_val)
        sql_b = build_custom_agg_sql(table_b, expressions, where_val)
        ra, rb = run_parallel(
            lambda: ex_a.run(sql_a, max_rows=1),
            lambda: ex_b.run(sql_b, max_rows=1),
        )
        row_a = ra.rows[0] if ra.rows else ()
        row_b = rb.rows[0] if rb.rows else ()
        result = compare_custom_aggs(table_a, row_a, table_b, row_b, expressions)
        return result, build_custom_agg_card(result, lang_val)

    def _compare_custom_agg_fn(table_a, table_b, lang_val, where_val, expressions_str):
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        if not expressions_str or not expressions_str.strip():
            return f'<div style="color:#dc2626;">{dc(lang_val, "dc_custom_agg_hint")}</div>'
        result, html = _custom_agg_inner(table_a, table_b, lang_val, where_val, expressions_str)
        _store_partial("custom_agg", result)
        return html

    _compare_custom_agg = _with_validation(holder, holder_lock, _compare_custom_agg_fn)

    # ── Custom SQL comparison (B) ──

    def _compare_sql(sql_a_text, sql_b_text, lang_val):
        with holder_lock:
            if holder.get("executor_a") is None or holder.get("executor_b") is None:
                return dc(lang_val, "dc_connect_both")
        if not sql_a_text.strip() or not sql_b_text.strip():
            return dc(lang_val, "dc_sql_readonly_only")

        for sql_text in (sql_a_text, sql_b_text):
            stripped = sql_text.strip().upper()
            if not (stripped.startswith("SELECT") or stripped.startswith("WITH")):
                return dc(lang_val, "dc_sql_readonly_only")
            if _DDL_KEYWORDS.search(sql_text):
                return dc(lang_val, "dc_sql_readonly_only")

        try:
            with holder_lock:
                ex_a, ex_b = holder["executor_a"], holder["executor_b"]
            ra, rb = run_parallel(
                lambda: ex_a.run(sql_a_text.strip(), max_rows=200),
                lambda: ex_b.run(sql_b_text.strip(), max_rows=200),
            )
            diff = diff_results(ra.columns, ra.rows, rb.columns, rb.rows)
            cols = ra.columns or rb.columns
            return build_sample_diff_card(diff, lang_val, columns=cols)
        except Exception as e:
            return _error_html(lang_val, e)

    # ── Compare All (F — structured report, G — schema cache, H — progress) ──

    def _compare_all(table_a, table_b, lang_val, where_val, key_cols_str,
                     threshold_str="", strategy="TOP N", col_mapping_str="",
                     masking_on=True, webhook_url="", notify_on_fail=False,
                     skew_cols_str="", stratified_col="",
                     checksum_cols_str="", partition_col="", custom_agg_str="",
                     progress=gr.Progress()):
        with holder_lock:
            if holder.get("executor_a") is None or holder.get("executor_b") is None:
                return dc(lang_val, "dc_connect_both")
        if not table_a or not table_b:
            return dc(lang_val, "dc_select_tables")
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg

        threshold = parse_threshold(threshold_str)

        def _safe_progress(val, desc=""):
            try:
                progress(val, desc=desc)
            except Exception:
                pass

        try:
            t0 = time.perf_counter()
            ex_a, ex_b = _snap_executors()

            _safe_progress(0.1, dc(lang_val, "dc_running"))
            desc_a, desc_b = run_parallel(
                lambda: ex_a.describe_table(table_a),
                lambda: ex_b.describe_table(table_b),
            )

            _safe_progress(0.25, dc(lang_val, "dc_running"))
            schema_result, schema_html = _schema_inner(table_a, table_b, lang_val, where_val, desc_a, desc_b)

            _safe_progress(0.4, dc(lang_val, "dc_running"))
            count_result, _ = _count_inner(table_a, table_b, lang_val, where_val)
            count_html = build_count_card(count_result, lang_val, threshold)

            _safe_progress(0.55, dc(lang_val, "dc_running"))
            sample_result, sample_html = _sample_inner(
                table_a, table_b, lang_val, where_val, key_cols_str,
                strategy, col_mapping_str, masking_on, stratified_col)

            _safe_progress(0.7, dc(lang_val, "dc_running"))
            agg_result, agg_html = _agg_inner(table_a, table_b, lang_val, where_val, desc_a, desc_b)

            _safe_progress(0.8, dc(lang_val, "dc_running"))
            profile_result, profile_html = _profile_inner(table_a, table_b, lang_val, where_val)

            skew_result = None
            skew_html = ""
            if skew_cols_str and skew_cols_str.strip():
                _safe_progress(0.84, dc(lang_val, "dc_running"))
                skew_result, skew_html = _skew_inner(
                    table_a, table_b, lang_val, where_val, skew_cols_str)

            checksum_result = None
            checksum_html = ""
            if checksum_cols_str and checksum_cols_str.strip():
                _safe_progress(0.88, dc(lang_val, "dc_running"))
                checksum_result, checksum_html = _checksum_inner(
                    table_a, table_b, lang_val, where_val, checksum_cols_str)

            partition_result = None
            partition_html = ""
            if partition_col and partition_col.strip():
                _safe_progress(0.92, dc(lang_val, "dc_running"))
                partition_result, partition_html = _partition_inner(
                    table_a, table_b, lang_val, where_val, partition_col)

            custom_agg_result = None
            custom_agg_html = ""
            if custom_agg_str and custom_agg_str.strip():
                _safe_progress(0.96, dc(lang_val, "dc_running"))
                custom_agg_result, custom_agg_html = _custom_agg_inner(
                    table_a, table_b, lang_val, where_val, custom_agg_str)

            elapsed = int((time.perf_counter() - t0) * 1000)

            keyed_diff = sample_result if isinstance(sample_result, KeyedDiffResult) else None
            sample_for_report = None if keyed_diff else sample_result

            report = CompareReport(
                schema=schema_result,
                row_count=count_result,
                sample=sample_for_report,
                aggregate=agg_result,
                profile=profile_result,
                keyed_diff=keyed_diff,
                skew=skew_result,
                checksum=checksum_result,
                partition=partition_result,
                custom_agg=custom_agg_result,
                elapsed_ms=elapsed,
            )
            with holder_lock:
                holder["last_report"] = report

            has_diffs = (
                (schema_result and schema_result.has_changes)
                or (count_result and count_result.delta != 0)
                or (agg_result and agg_result.mismatches > 0)
                or (checksum_result and checksum_result.mismatch_count > 0)
                or (partition_result and partition_result.total_delta != 0)
                or (custom_agg_result and custom_agg_result.mismatches > 0)
            )

            if webhook_url and webhook_url.strip():
                should_send = True
                if notify_on_fail and not has_diffs:
                    should_send = False
                if should_send:
                    payload = {
                        "table_a": table_a, "table_b": table_b,
                        "has_diffs": has_diffs,
                        "elapsed_ms": elapsed,
                        "row_count_delta": count_result.delta if count_result else 0,
                        "schema_changes": len(schema_result.items) if schema_result and schema_result.has_changes else 0,
                        "aggregate_mismatches": agg_result.mismatches if agg_result else 0,
                        "checksum_mismatches": checksum_result.mismatch_count if checksum_result else 0,
                        "timestamp": datetime.now(timezone.utc).isoformat(),
                        "summary": {
                            "count_a": count_result.count_a if count_result else 0,
                            "count_b": count_result.count_b if count_result else 0,
                            "delta_pct": count_result.delta_pct if count_result else 0,
                        },
                        "report_summary": _build_webhook_summary(report),
                    }
                    threading.Thread(
                        target=_send_webhook,
                        args=(webhook_url.strip(), payload),
                        daemon=True,
                    ).start()

            _safe_progress(1.0, "Done")
            summary_html = build_summary_card(report, lang_val)
            return (summary_html + schema_html + count_html + sample_html
                    + agg_html + profile_html + skew_html
                    + checksum_html + partition_html + custom_agg_html)

        except Exception as e:
            if webhook_url and webhook_url.strip():
                threading.Thread(
                    target=_send_webhook,
                    args=(webhook_url.strip(), {"error": str(e), "table_a": table_a, "table_b": table_b}),
                    daemon=True,
                ).start()
            return _error_html(lang_val, e)

    # ── Batch count (H — progress) ──

    def _batch_count(lang_val, where_val, progress=gr.Progress()):
        with holder_lock:
            if holder.get("executor_a") is None or holder.get("executor_b") is None:
                return dc(lang_val, "dc_connect_both")
            tables_a = list(holder["tables_a"])
            tables_b = list(holder["tables_b"])
            ex_a, ex_b = holder["executor_a"], holder["executor_b"]

        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg

        try:
            pairs = find_common_tables(tables_a, tables_b)
            if not pairs:
                return (
                    '<div style="border:1px solid #e5e7eb;border-radius:8px;padding:12px;'
                    'background:#f9fafb;margin-bottom:8px;">'
                    f'<b style="color:#6b7280;">{dc(lang_val, "dc_batch_result")}</b> — '
                    f'{dc(lang_val, "dc_no_common")}</div>'
                )

            results: list[RowCountResult] = []
            total = len(pairs)
            for idx, (ta, tb) in enumerate(pairs):
                progress((idx + 1) / total, desc=dc(lang_val, "dc_comparing_table").format(name=ta))
                sql_a = build_count_sql(ta, where_val)
                sql_b = build_count_sql(tb, where_val)
                ra, rb = run_parallel(
                    lambda: ex_a.run(sql_a, max_rows=1),
                    lambda: ex_b.run(sql_b, max_rows=1),
                )
                ca = int(ra.rows[0][0]) if ra.rows else 0
                cb = int(rb.rows[0][0]) if rb.rows else 0
                results.append(build_row_count_result(ta, ca, tb, cb))

            _store_partial("batch_counts", results)

            return build_batch_count_card(results, lang_val)
        except Exception as e:
            return _error_html(lang_val, e)

    # ── Report persistence (E) ──

    def _save_report(lang_val):
        with holder_lock:
            report = holder.get("last_report")
        if not report:
            return dc(lang_val, "dc_error")
        try:
            _REPORTS_DIR.mkdir(parents=True, exist_ok=True)
            ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
            fpath = _REPORTS_DIR / f"compare_{ts}.json"
            data = report.to_dict()
            with open(fpath, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            return f'✅ {dc(lang_val, "dc_saved")} — {fpath.name}'
        except Exception as e:
            return f'❌ {dc(lang_val, "dc_error")}: {_esc_html(str(e))}'

    def _list_reports():
        if not _REPORTS_DIR.is_dir():
            return gr.update(choices=[], value=None)
        files = sorted(_REPORTS_DIR.glob("compare_*.json"), reverse=True)
        names = [f.name for f in files[:_MAX_REPORT_FILES]]
        return gr.update(choices=names, value=None)

    def _load_report(filename, lang_val):
        if not filename:
            return dc(lang_val, "dc_no_reports")
        fpath = _REPORTS_DIR / filename
        if not fpath.is_file():
            return dc(lang_val, "dc_no_reports")
        try:
            with open(fpath, "r", encoding="utf-8") as f:
                data = json.load(f)
            report = CompareReport.from_dict(data)
            with holder_lock:
                holder["last_report"] = report

            parts: list[str] = []
            parts.append(f'<div style="color:#16a34a;font-weight:600;margin-bottom:8px;">'
                         f'✅ {dc(lang_val, "dc_loaded")} — {_esc_html(filename)}</div>')
            if report.schema is not None:
                parts.append(build_schema_diff_card(report.schema, lang_val))
            if report.row_count is not None:
                parts.append(build_count_card(report.row_count, lang_val))
            if report.aggregate is not None:
                parts.append(build_aggregate_card(report.aggregate, lang_val))
            if report.profile is not None:
                parts.append(build_profile_card(report.profile, lang_val))
            if report.skew is not None:
                parts.append(build_skew_card(report.skew, lang_val))
            if report.checksum is not None:
                parts.append(build_checksum_card(report.checksum, lang_val))
            if report.partition is not None:
                parts.append(build_partition_card(report.partition, lang_val))
            if report.custom_agg is not None:
                parts.append(build_custom_agg_card(report.custom_agg, lang_val))
            if report.keyed_diff is not None:
                parts.append(build_keyed_diff_card(report.keyed_diff, lang_val))
            if report.batch_counts is not None:
                parts.append(build_batch_count_card(report.batch_counts, lang_val))
            return "\n".join(parts) if parts else dc(lang_val, "dc_no_reports")
        except Exception as e:
            return _error_html(lang_val, e)

    # ── Export ──

    def _export_csv(lang_val):
        with holder_lock:
            report = holder.get("last_report")
        if not report:
            return None
        try:
            return _export_report_csv(report)
        except Exception:
            _log.warning("CSV export failed", exc_info=True)
            return None

    def _export_excel(lang_val):
        with holder_lock:
            report = holder.get("last_report")
        if not report:
            return None
        try:
            return _export_report_excel(report)
        except Exception:
            _log.warning("Excel export failed", exc_info=True)
            return None

    # ── Connection presets (A) ──

    def _refresh_presets():
        names = [d["name"] for d in _PRESETS_STORE.list()]
        return gr.update(choices=names, value=None), gr.update(choices=names, value=None)

    def _save_preset(name, ds_label, host, port, db, user, pwd, env_val, lang_val):
        # Hidden auth/host textboxes submit None from the browser, not "".
        name = (name or "").strip()
        if not name:
            return dc(lang_val, "dc_error")
        ds_type = _DS_LABEL_TO_KEY.get(ds_label)
        if ds_type is None:
            return f"❌ {dc(lang_val, 'dc_error')}: unknown datasource {_esc_html(str(ds_label))}"
        try:
            port_int = int(port) if port else 0
        except ValueError:
            port_int = 0
        try:
            _PRESETS_STORE.save(name, ds_type, (host or "").strip(), port_int,
                                (db or "").strip(), (user or "").strip(),
                                (pwd or "").strip(),
                                environment=env_val.strip() if env_val else "")
            return f"✅ {dc(lang_val, 'dc_preset_saved')}: {_esc_html(name)}"
        except Exception as e:
            return _error_html(lang_val, e)

    def _load_preset(name, lang_val):
        if not name:
            return [gr.update()] * 7
        try:
            preset = _PRESETS_STORE.get_by_name(name)
            if not preset:
                return [gr.update()] * 7
            dialect = DIALECT_NAMES.get(preset["ds_type"], preset["ds_type"])
            return (
                gr.update(value=dialect),
                gr.update(value=preset.get("host", "")),
                gr.update(value=str(preset.get("port", ""))),
                gr.update(value=preset.get("database", "")),
                gr.update(value=preset.get("username", "")),
                gr.update(value=preset.get("password", "")),
                gr.update(value=preset.get("environment", "")),
            )
        except Exception:
            _log.warning("Failed to load preset: %s", name, exc_info=True)
            return [gr.update()] * 7

    def _delete_preset(name, lang_val):
        if not name:
            return dc(lang_val, "dc_preset_no_presets")
        try:
            presets = _PRESETS_STORE.list()
            for p in presets:
                if p["name"] == name:
                    _PRESETS_STORE.delete(p["id"])
                    return f"✅ {dc(lang_val, 'dc_preset_deleted')}"
            return dc(lang_val, "dc_preset_no_presets")
        except Exception as e:
            return _error_html(lang_val, e)

    # ── Batch full comparison (F) ──

    def _batch_full(lang_val, where_val, threshold_str, progress=gr.Progress()):
        with holder_lock:
            if holder.get("executor_a") is None or holder.get("executor_b") is None:
                return dc(lang_val, "dc_connect_both")
            tables_a = list(holder["tables_a"])
            tables_b = list(holder["tables_b"])
            ex_a, ex_b = holder["executor_a"], holder["executor_b"]

        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg

        threshold = parse_threshold(threshold_str)

        try:
            pairs = find_common_tables(tables_a, tables_b)
            if not pairs:
                return (
                    '<div style="border:1px solid #e5e7eb;border-radius:8px;padding:12px;'
                    'background:#f9fafb;margin-bottom:8px;">'
                    f'<b style="color:#6b7280;">{dc(lang_val, "dc_batch_full_result")}</b> — '
                    f'{dc(lang_val, "dc_no_common")}</div>'
                )

            items: list[BatchFullItem] = []
            total = len(pairs)
            for idx, (ta, tb) in enumerate(pairs):
                progress((idx + 1) / total,
                         desc=dc(lang_val, "dc_comparing_table").format(name=ta))
                try:
                    schema_r, _ = _schema_inner(ta, tb, lang_val, where_val)
                except Exception:
                    schema_r = None
                try:
                    count_r, _ = _count_inner(ta, tb, lang_val, where_val)
                except Exception:
                    count_r = None
                try:
                    agg_r, _ = _agg_inner(ta, tb, lang_val, where_val)
                except Exception:
                    agg_r = None

                has_diffs = False
                if schema_r and schema_r.has_changes:
                    has_diffs = True
                if count_r and count_r.delta != 0:
                    has_diffs = True
                if agg_r and agg_r.mismatches > 0:
                    has_diffs = True

                items.append(BatchFullItem(ta, tb, schema_r, count_r, agg_r, has_diffs))

            _store_partial("batch_full", items)
            return build_batch_full_card(items, lang_val, threshold)
        except Exception as e:
            return _error_html(lang_val, e)

    # ── SeaTunnel sync config (C) ──

    def _gen_sync_config(table_a, table_b, col_mapping_str, lang_val):
        with holder_lock:
            ex_a = holder.get("executor_a")
            ex_b = holder.get("executor_b")
        if not ex_a or not ex_b:
            return dc(lang_val, "dc_connect_both")
        if not table_a or not table_b:
            return dc(lang_val, "dc_select_tables")
        try:
            cfg_a, cfg_b = ex_a.config, ex_b.config
            mapping = parse_column_mapping(col_mapping_str) if col_mapping_str else None
            config_text = generate_sync_config(
                cfg_a.ds_type, cfg_a.host, cfg_a.port, cfg_a.database,
                cfg_a.username, cfg_a.password, table_a,
                cfg_b.ds_type, cfg_b.host, cfg_b.port, cfg_b.database,
                cfg_b.username, cfg_b.password, table_b,
                column_mapping=mapping,
            )
            header = f'<div style="color:#16a34a;font-weight:600;margin-bottom:8px;">✅ {dc(lang_val, "dc_sync_generated")}</div>'
            code_block = f'<pre style="background:#1e293b;color:#e2e8f0;padding:16px;border-radius:8px;overflow-x:auto;font-size:12px;">{_esc_html(config_text)}</pre>'
            return header + code_block
        except Exception as e:
            return _error_html(lang_val, e)

    # ── Scheduled comparison (E) ──

    def _schedule_tick(table_a, table_b, lang_val, where_val, key_cols_str, interval_min):
        """Background timer callback — run comparison and reschedule."""
        try:
            _compare_all(table_a, table_b, lang_val, where_val, key_cols_str)
            _save_report(lang_val)
            with holder_lock:
                holder["schedule_last_run"] = datetime.now(timezone.utc).strftime("%H:%M:%S")
        except Exception:
            _log.warning("Scheduled comparison failed", exc_info=True)
        with holder_lock:
            if holder["schedule_active"]:
                t = threading.Timer(
                    float(interval_min) * 60,
                    _schedule_tick,
                    args=(table_a, table_b, lang_val, where_val, key_cols_str, interval_min),
                )
                t.daemon = True
                holder["schedule_timer"] = t
                t.start()

    def _schedule_start(table_a, table_b, lang_val, where_val, key_cols_str, interval_str):
        if not table_a or not table_b:
            return dc(lang_val, "dc_select_tables")
        try:
            interval_min = int(interval_str)
        except (ValueError, TypeError):
            interval_min = _DEFAULT_SCHEDULE_MIN
        interval_min = max(1, interval_min)
        with holder_lock:
            if holder["schedule_active"]:
                return dc(lang_val, "dc_schedule_running").format(m=interval_min)
            holder["schedule_active"] = True
        t = threading.Timer(
            float(interval_min) * 60,
            _schedule_tick,
            args=(table_a, table_b, lang_val, where_val, key_cols_str, interval_min),
        )
        t.daemon = True
        with holder_lock:
            holder["schedule_timer"] = t
        t.start()
        return f'✅ {dc(lang_val, "dc_schedule_running").format(m=interval_min)}'

    def _schedule_stop(lang_val):
        with holder_lock:
            holder["schedule_active"] = False
            timer = holder.get("schedule_timer")
            if timer:
                timer.cancel()
                holder["schedule_timer"] = None
        return dc(lang_val, "dc_schedule_stopped")

    # ── History trends (H) ──

    def _show_trend(lang_val, alert_rules_raw=""):
        entries = build_trend_data(_REPORTS_DIR)
        if not entries:
            return dc(lang_val, "dc_trend_no_data"), gr.update(visible=False)
        try:
            plt = _get_matplotlib()
            fig, axes = plt.subplots(3, 1, figsize=(8, 9), sharex=True)
            timestamps = [e["timestamp"] for e in entries]
            x = range(len(timestamps))

            # Subplot 1: Row Count Delta (line)
            deltas = [e["delta"] for e in entries]
            axes[0].plot(x, deltas, marker="o", color="#4C78A8",
                         linewidth=2, markersize=5)
            axes[0].set_ylabel("Row Count Delta")
            axes[0].set_title(dc(lang_val, "dc_trend_chart"))
            axes[0].axhline(y=0, color="#94a3b8", linestyle="--", linewidth=0.8)
            axes[0].grid(True, alpha=0.3)

            # Subplot 2: Aggregate Mismatches (bar)
            mismatches = [e.get("mismatches", 0) for e in entries]
            axes[1].bar(x, mismatches, color="#f59e0b", alpha=0.8)
            axes[1].set_ylabel(dc(lang_val, "dc_trend_mismatches"))
            axes[1].grid(True, alpha=0.3)

            # Subplot 3: Schema Changes (bar)
            schema_changes = [e.get("schema_changes", 0) for e in entries]
            axes[2].bar(x, schema_changes, color="#ef4444", alpha=0.8)
            axes[2].set_ylabel(dc(lang_val, "dc_trend_schema"))
            axes[2].set_xticks(list(x))
            axes[2].set_xticklabels(timestamps, rotation=45, ha="right", fontsize=8)
            axes[2].grid(True, alpha=0.3)

            fig.set_facecolor("#fafafa")
            fig.tight_layout()
            info = dc(lang_val, "dc_history_count").format(n=len(entries))

            if alert_rules_raw and alert_rules_raw.strip():
                rules = parse_alert_rules(alert_rules_raw)
                if rules:
                    alerts = check_trend_alerts(entries, rules)
                    info += "<br>" + build_trend_alert_card(alerts, lang_val)

            return info, gr.update(visible=True, value=fig)
        except Exception as e:
            return f'❌ {_esc_html(str(e))}', gr.update(visible=False)

    # ── Report diff (Phase 5A) ──

    def _compare_reports_fn(old_name, new_name, lang_val):
        if not old_name or not new_name:
            return dc(lang_val, "dc_report_diff")
        try:
            old_path = _REPORTS_DIR / old_name
            new_path = _REPORTS_DIR / new_name
            if not old_path.is_file() or not new_path.is_file():
                return dc(lang_val, "dc_no_reports")
            with open(old_path, "r", encoding="utf-8") as f:
                old_report = CompareReport.from_dict(json.load(f))
            with open(new_path, "r", encoding="utf-8") as f:
                new_report = CompareReport.from_dict(json.load(f))
            rdiff = diff_reports(old_report, new_report)
            return build_report_diff_card(rdiff, lang_val)
        except Exception as e:
            return _error_html(lang_val, e)

    # ── Data lineage (Phase 5C) ──

    def _show_lineage_fn(table_a, lang_val, lineage_sql_str):
        if not table_a:
            return dc(lang_val, "dc_select_tables")
        if not lineage_sql_str or not lineage_sql_str.strip():
            return dc(lang_val, "dc_lineage_hint")
        try:
            sqls = [s.strip() for s in lineage_sql_str.strip().split("\n") if s.strip()]
            upstream = get_upstream_tables(table_a, sqls)
            return build_lineage_card(table_a, upstream, lang_val)
        except Exception as e:
            return _error_html(lang_val, e)

    # ── Feature A — Incremental comparison ──

    def _incremental_fn(table_a, table_b, lang_val, where_val,
                        watermark_col, watermark_val, key_cols_str,
                        col_mapping_str, masking_on):
        with holder_lock:
            if holder.get("executor_a") is None or holder.get("executor_b") is None:
                return dc(lang_val, "dc_connect_both")
        if not table_a or not table_b:
            return dc(lang_val, "dc_select_tables")
        if not watermark_col.strip():
            return dc(lang_val, "dc_watermark_hint")
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        try:
            ex_a, ex_b = holder["executor_a"], holder["executor_b"]
            last_val = watermark_val.strip()
            sql_a = build_incremental_sql(table_a, watermark_col.strip(), last_val, 1000, where_val,
                                         ds_type=ex_a.config.ds_type)
            sql_b = build_incremental_sql(table_b, watermark_col.strip(), last_val, 1000, where_val,
                                         ds_type=ex_b.config.ds_type)
            ra, rb = run_parallel(
                lambda: ex_a.run(sql_a, max_rows=1000),
                lambda: ex_b.run(sql_b, max_rows=1000),
            )
            mapping = parse_column_mapping(col_mapping_str) if col_mapping_str else {}
            rb_cols = apply_column_mapping(rb.columns, mapping) if mapping else rb.columns
            rows_a, rows_b = list(ra.rows), list(rb.rows)
            if masking_on:
                sa = detect_sensitive_columns(ra.columns, rows_a)
                sb = detect_sensitive_columns(rb_cols, rows_b)
                if sa:
                    rows_a = mask_rows(ra.columns, rows_a, sa)
                if sb:
                    rows_b = mask_rows(rb_cols, rows_b, sb)
            key_cols = [k.strip() for k in key_cols_str.split(",") if k.strip()] if key_cols_str else []
            if key_cols:
                kd_result = diff_by_key(ra.columns, rows_a, rb_cols, rows_b, key_cols)
                _store_partial("keyed_diff", kd_result)
                return build_keyed_diff_card(kd_result, lang_val)
            diff = diff_results(ra.columns, rows_a, rb_cols, rows_b)
            _store_partial("sample", diff)
            return build_sample_diff_card(diff, lang_val, columns=ra.columns or rb_cols)
        except Exception as e:
            return _error_html(lang_val, e)

    # ── Feature B — Quality rules ──

    def _check_quality_fn(table_a, table_b, lang_val, where_val, quality_rules_str):
        with holder_lock:
            if holder.get("executor_a") is None or holder.get("executor_b") is None:
                return dc(lang_val, "dc_connect_both")
        if not table_a or not table_b:
            return dc(lang_val, "dc_select_tables")
        if not quality_rules_str.strip():
            return dc(lang_val, "dc_quality_rules_hint")
        try:
            rules = parse_quality_rules(quality_rules_str)
            if not rules:
                return dc(lang_val, "dc_quality_rules_hint")
            profile_result, _ = _profile_inner(table_a, table_b, lang_val, where_val)
            results = check_quality_rules(rules, profile_result)
            return build_quality_card(results, lang_val)
        except Exception as e:
            return _error_html(lang_val, e)

    # ── Feature C — Diff SQL generation ──

    def _gen_diff_sql_fn(table_a, table_b, lang_val):
        with holder_lock:
            report = holder.get("last_report")
        if not report or not report.keyed_diff:
            return dc(lang_val, "dc_diff_sql_no_diff")
        try:
            with holder_lock:
                ex_b = holder.get("executor_b")
            ds_type = ex_b.config.ds_type if ex_b else "mysql"
            sql = generate_diff_sql(report.keyed_diff, table_b, ds_type)
            return build_diff_sql_card(sql, lang_val)
        except Exception as e:
            return _error_html(lang_val, e)

    # ── Feature D — Templates ──

    def _refresh_templates():
        names = [d["name"] for d in _TEMPLATES_STORE.list()]
        return gr.update(choices=names, value=None)

    def _save_template(name, table_a, table_b, where_val, key_cols_str,
                       threshold_str, col_mapping_str, strategy, watermark_col,
                       quality_rules_str, webhook_url, notify_on_fail, masking_on, lang_val):
        if not name.strip():
            return dc(lang_val, "dc_error"), gr.update()
        try:
            _TEMPLATES_STORE.save(
                name=name.strip(), table_a=table_a or "", table_b=table_b or "",
                where_clause=where_val, key_columns=key_cols_str,
                threshold=threshold_str, column_mapping=col_mapping_str,
                sample_strategy=strategy, watermark_column=watermark_col,
                quality_rules=quality_rules_str, webhook_url=webhook_url,
                notify_on_fail=bool(notify_on_fail), masking_enabled=bool(masking_on),
            )
            names = [d["name"] for d in _TEMPLATES_STORE.list()]
            return f"✅ {dc(lang_val, 'dc_template_saved')}", gr.update(choices=names, value=None)
        except Exception as e:
            return _error_html(lang_val, e), gr.update()

    def _load_template(name, lang_val):
        if not name:
            return [gr.update()] * 12
        try:
            tmpl = _TEMPLATES_STORE.get_by_name(name)
            if not tmpl:
                return [gr.update()] * 12
            return (
                gr.update(value=tmpl.get("table_a", "")),
                gr.update(value=tmpl.get("table_b", "")),
                gr.update(value=tmpl.get("where_clause", "")),
                gr.update(value=tmpl.get("key_columns", "")),
                gr.update(value=tmpl.get("threshold", "")),
                gr.update(value=tmpl.get("column_mapping", "")),
                gr.update(value=tmpl.get("sample_strategy", "TOP N")),
                gr.update(value=tmpl.get("watermark_column", "")),
                gr.update(value=tmpl.get("quality_rules", "")),
                gr.update(value=tmpl.get("webhook_url", "")),
                gr.update(value=tmpl.get("notify_on_fail", False)),
                gr.update(value=tmpl.get("masking_enabled", True)),
            )
        except Exception:
            _log.warning("Failed to load template", exc_info=True)
            return [gr.update()] * 12

    def _delete_template(name, lang_val):
        if not name:
            return dc(lang_val, "dc_template_no_templates"), gr.update()
        try:
            for t in _TEMPLATES_STORE.list():
                if t["name"] == name:
                    _TEMPLATES_STORE.delete(t["id"])
                    names = [d["name"] for d in _TEMPLATES_STORE.list()]
                    return f"✅ {dc(lang_val, 'dc_template_deleted')}", gr.update(choices=names, value=None)
            return dc(lang_val, "dc_template_no_templates"), gr.update()
        except Exception as e:
            return _error_html(lang_val, e), gr.update()

    # ── Feature F6 — Batch template orchestration ──

    def _batch_templates_fn(selected_names, lang_val):
        if not selected_names:
            return dc(lang_val, "dc_batch_templates_hint")
        with holder_lock:
            ex_a = holder.get("executor_a")
            ex_b = holder.get("executor_b")
        if ex_a is None or ex_b is None:
            return dc(lang_val, "dc_connect_both")
        results = []
        for name in selected_names:
            tmpl = _TEMPLATES_STORE.get_by_name(name)
            if not tmpl:
                results.append({"name": name, "status": "skip", "detail": "not found"})
                continue
            try:
                ta = tmpl.get("table_a", "")
                tb = tmpl.get("table_b", "")
                wh = tmpl.get("where_clause", "")

                # Count comparison
                count_sql_a = build_count_sql(ta, wh)
                count_sql_b = build_count_sql(tb, wh)
                ra, rb = run_parallel(
                    lambda s=count_sql_a: ex_a.run(s, max_rows=1),
                    lambda s=count_sql_b: ex_b.run(s, max_rows=1),
                )
                cnt_a = int(ra.rows[0][0]) if ra.rows else 0
                cnt_b = int(rb.rows[0][0]) if rb.rows else 0

                # Schema comparison
                schema_changes = 0
                try:
                    desc_a, desc_b = run_parallel(
                        lambda t=ta: ex_a.describe_table(t),
                        lambda t=tb: ex_b.describe_table(t),
                    )
                    ca_list = [{"name": c.name, "type": c.dtype} for c in desc_a.columns]
                    cb_list = [{"name": c.name, "type": c.dtype} for c in desc_b.columns]
                    schema_result = compare_schemas(ca_list, cb_list, ta, tb)
                    schema_changes = len(schema_result.items) if schema_result.has_changes else 0
                except Exception:
                    _log.warning("Batch template %r: schema comparison failed",
                                 name, exc_info=True)
                    desc_a = desc_b = None

                # Aggregate comparison on shared numeric columns
                agg_mismatches = 0
                try:
                    if desc_a and desc_b:
                        names_b = {c.name.lower() for c in desc_b.columns}
                        shared_num = [
                            c.name for c in desc_a.columns
                            if is_numeric_type(c.dtype) and c.name.lower() in names_b
                        ]
                        if shared_num:
                            agg_sql_a = build_aggregate_sql(
                                ta, shared_num[:_MAX_AGG_COLUMNS], wh,
                                ds_type=ex_a.config.ds_type)
                            agg_sql_b = build_aggregate_sql(
                                tb, shared_num[:_MAX_AGG_COLUMNS], wh,
                                ds_type=ex_b.config.ds_type)
                            agg_ra, agg_rb = run_parallel(
                                lambda s=agg_sql_a: ex_a.run(s, max_rows=200),
                                lambda s=agg_sql_b: ex_b.run(s, max_rows=200),
                            )
                            agg_result = compare_aggregates(
                                ta, agg_ra.rows[0] if agg_ra.rows else (),
                                tb, agg_rb.rows[0] if agg_rb.rows else (),
                                shared_num[:_MAX_AGG_COLUMNS],
                            )
                            agg_mismatches = agg_result.mismatches
                except Exception:
                    _log.warning("Batch template %r: aggregate comparison failed",
                                 name, exc_info=True)

                passed = cnt_a == cnt_b and schema_changes == 0 and agg_mismatches == 0
                results.append({
                    "name": name, "table_a": ta, "table_b": tb,
                    "count_a": cnt_a, "count_b": cnt_b,
                    "schema_changes": schema_changes,
                    "agg_mismatches": agg_mismatches,
                    "status": "pass" if passed else "fail",
                })
            except Exception as exc:
                results.append({"name": name, "status": "fail", "detail": str(exc)})
        return build_batch_template_card(results, lang_val)

    # ── Feature H — Multi-environment ──

    def _compare_environments_fn(env_a, env_b, table_name, lang_val, where_val):
        if not env_a.strip() or not env_b.strip():
            return dc(lang_val, "dc_env_hint")
        if not table_name.strip():
            return dc(lang_val, "dc_select_tables")
        ok, msg = _validate_where(where_val, lang_val)
        if not ok:
            return msg
        presets = _PRESETS_STORE.list()
        pa = next((p for p in presets if p.get("environment", "") == env_a.strip()), None)
        pb = next((p for p in presets if p.get("environment", "") == env_b.strip()), None)
        if not pa or not pb:
            return dc(lang_val, "dc_env_hint")
        try:
            cfg_a = DatabaseConfig(
                ds_type=pa["ds_type"], host=pa["host"], port=pa["port"],
                database=pa["database"], username=pa.get("username"),
                password=pa.get("password"),
            )
            cfg_b = DatabaseConfig(
                ds_type=pb["ds_type"], host=pb["host"], port=pb["port"],
                database=pb["database"], username=pb.get("username"),
                password=pb.get("password"),
            )
            ex_a = create_executor(cfg_a)
            ex_b = create_executor(cfg_b)
            tbl = table_name.strip()
            cards: list[str] = []
            label_a = f"{tbl} ({env_a})"
            label_b = f"{tbl} ({env_b})"

            # Schema comparison
            desc_a, desc_b = run_parallel(
                lambda: ex_a.describe_table(tbl),
                lambda: ex_b.describe_table(tbl),
            )
            ca_list = [{"name": c.name, "type": c.dtype} for c in desc_a.columns]
            cb_list = [{"name": c.name, "type": c.dtype} for c in desc_b.columns]
            schema_result = compare_schemas(ca_list, cb_list, label_a, label_b)
            cards.append(build_schema_diff_card(schema_result, lang_val))

            # Count comparison
            sql_ca = build_count_sql(tbl, where_val)
            sql_cb = build_count_sql(tbl, where_val)
            ra, rb = run_parallel(
                lambda: ex_a.run(sql_ca, max_rows=1),
                lambda: ex_b.run(sql_cb, max_rows=1),
            )
            ca = int(ra.rows[0][0]) if ra.rows else 0
            cb = int(rb.rows[0][0]) if rb.rows else 0
            count_result = build_row_count_result(label_a, ca, label_b, cb)
            cards.append(build_count_card(count_result, lang_val))

            # Aggregate comparison on shared numeric columns
            names_b = {c.name.lower() for c in desc_b.columns}
            shared_num = [
                c.name for c in desc_a.columns
                if is_numeric_type(c.dtype) and c.name.lower() in names_b
            ]
            if shared_num:
                agg_sql_a = build_aggregate_sql(tbl, shared_num[:_MAX_AGG_COLUMNS],
                                                where_val, ds_type=ex_a.config.ds_type)
                agg_sql_b = build_aggregate_sql(tbl, shared_num[:_MAX_AGG_COLUMNS],
                                                where_val, ds_type=ex_b.config.ds_type)
                agg_ra, agg_rb = run_parallel(
                    lambda: ex_a.run(agg_sql_a, max_rows=200),
                    lambda: ex_b.run(agg_sql_b, max_rows=200),
                )
                agg_result = compare_aggregates(
                    label_a, agg_ra.rows[0] if agg_ra.rows else (),
                    label_b, agg_rb.rows[0] if agg_rb.rows else (),
                    shared_num[:_MAX_AGG_COLUMNS],
                )
                cards.append(build_aggregate_card(agg_result, lang_val))

            header = f'<h4 style="color:#6366f1;">{dc(lang_val, "dc_env_compare_result")}</h4>'
            return header + "\n".join(cards)
        except Exception as e:
            return _error_html(lang_val, e)

    # ── DS type change handler ──
    def _on_ds_change(ds_label: str):
        ds = _DS_LABEL_TO_KEY.get(ds_label)
        if ds is None:
            return tuple(gr.update() for _ in range(5))
        defaults = DS_DEFAULTS[ds]
        default_port = str(defaults["port"])
        default_db = defaults["database"]
        show_auth = ds in _NEEDS_AUTH
        show_host = ds in _NEEDS_HOST
        env_cfg = config_from_env(ds)
        if env_cfg:
            return (
                gr.update(value=str(env_cfg.port), placeholder=default_port, visible=show_host),
                gr.update(value=env_cfg.database, placeholder=str(default_db)),
                gr.update(value=env_cfg.host, visible=show_host),
                gr.update(value=env_cfg.username or "", visible=show_auth),
                gr.update(value=env_cfg.password or "", visible=show_auth),
            )
        return (
            gr.update(value="", placeholder=default_port, visible=show_host),
            gr.update(value="", placeholder=str(default_db)),
            gr.update(value="", visible=show_host),
            gr.update(value="", visible=show_auth),
            gr.update(value="", visible=show_auth),
        )

    # ── Layout ──

    with gr.Row(elem_classes=["st-page-row"]):
        lang_state = gr.State("en")

        # ── Left sidebar ──
        with gr.Column(scale=0, min_width=720,
                       elem_classes=["st-sidebar", "st-dc-sidebar"]):
            title_md = gr.Markdown(t("dc_title"))

            with gr.Row():
                # Source A panel
                with gr.Column(min_width=260):
                    src_a_md = gr.Markdown(f"**{t('dc_source_a')}**")
                    ds_a = gr.Dropdown(choices=_DS_CHOICES, value=_DS_CHOICES[0],
                                       label=t("dc_datasource_type"), elem_classes=["st-sidebar-control"])
                    host_a = gr.Textbox(label=t("dc_host"), placeholder="10.0.0.1",
                                        elem_classes=["st-sidebar-control"])
                    with gr.Row():
                        port_a = gr.Textbox(label=t("dc_port"), placeholder="10000",
                                            elem_classes=["st-sidebar-control"])
                        db_a = gr.Textbox(label=t("dc_database"), placeholder="default",
                                          elem_classes=["st-sidebar-control"])
                    user_a = gr.Textbox(label=t("dc_username"), visible=False,
                                        elem_classes=["st-sidebar-control"])
                    pwd_a = gr.Textbox(label=t("dc_password"), type="password", visible=False,
                                       elem_classes=["st-sidebar-control"])
                    conn_a = gr.Button(t("dc_connect"), variant="secondary", size="sm",
                                       elem_classes=["st-connect-btn"])
                    status_a = gr.Textbox(label=t("dc_status"), value=t("dc_not_connected"),
                                          interactive=False, elem_classes=["st-sidebar-status"])
                    search_a = gr.Textbox(label="", placeholder=t("dc_search_tables"),
                                          elem_classes=["st-sidebar-control"])
                    table_a = gr.Dropdown(choices=[], label=t("dc_select_table"),
                                          elem_classes=["st-sidebar-control"])

                # Source B panel
                with gr.Column(min_width=260):
                    src_b_md = gr.Markdown(f"**{t('dc_source_b')}**")
                    ds_b = gr.Dropdown(choices=_DS_CHOICES, value=_DS_CHOICES[0],
                                       label=t("dc_datasource_type"), elem_classes=["st-sidebar-control"])
                    host_b = gr.Textbox(label=t("dc_host"), placeholder="10.0.0.1",
                                        elem_classes=["st-sidebar-control"])
                    with gr.Row():
                        port_b = gr.Textbox(label=t("dc_port"), placeholder="10000",
                                            elem_classes=["st-sidebar-control"])
                        db_b = gr.Textbox(label=t("dc_database"), placeholder="default",
                                          elem_classes=["st-sidebar-control"])
                    user_b = gr.Textbox(label=t("dc_username"), visible=False,
                                        elem_classes=["st-sidebar-control"])
                    pwd_b = gr.Textbox(label=t("dc_password"), type="password", visible=False,
                                       elem_classes=["st-sidebar-control"])
                    conn_b = gr.Button(t("dc_connect"), variant="secondary", size="sm",
                                       elem_classes=["st-connect-btn"])
                    status_b = gr.Textbox(label=t("dc_status"), value=t("dc_not_connected"),
                                          interactive=False, elem_classes=["st-sidebar-status"])
                    search_b = gr.Textbox(label="", placeholder=t("dc_search_tables"),
                                          elem_classes=["st-sidebar-control"])
                    table_b = gr.Dropdown(choices=[], label=t("dc_select_table"),
                                          elem_classes=["st-sidebar-control"])

            # Connection presets (A)
            with gr.Accordion(t("dc_presets"), open=False) as presets_accordion:
                with gr.Row():
                    preset_name_input = gr.Textbox(label=t("dc_preset_name"), scale=2,
                                                    elem_classes=["st-sidebar-control"])
                    preset_env_input = gr.Textbox(label=t("dc_preset_env"), scale=1,
                                                   elem_classes=["st-sidebar-control"])
                    preset_save_btn = gr.Button(f'{t("dc_preset_save")} (A)', size="sm",
                                                 elem_classes=["st-connect-btn"])
                    preset_save_b_btn = gr.Button(f'{t("dc_preset_save")} (B)', size="sm",
                                                   elem_classes=["st-connect-btn"])
                with gr.Row():
                    preset_dd_a = gr.Dropdown(choices=[], label=f"{t('dc_preset_load')} (A)",
                                               elem_classes=["st-sidebar-control"])
                    preset_dd_b = gr.Dropdown(choices=[], label=f"{t('dc_preset_load')} (B)",
                                               elem_classes=["st-sidebar-control"])
                preset_status = gr.HTML("")

            # D — WHERE filter
            where_input = gr.Textbox(label=t("dc_where_clause"),
                                      placeholder=t("dc_where_hint"),
                                      elem_classes=["st-sidebar-control"])

            # B — Threshold
            threshold_input = gr.Textbox(label=t("dc_threshold"),
                                          placeholder=t("dc_threshold_hint"),
                                          elem_classes=["st-sidebar-control"])

            # A — Key columns
            key_input = gr.Textbox(label=t("dc_key_columns"),
                                    placeholder="id, order_id",
                                    elem_classes=["st-sidebar-control"])

            # D — Column mapping
            mapping_input = gr.Textbox(label=t("dc_col_mapping"),
                                        placeholder=t("dc_col_mapping_hint"),
                                        elem_classes=["st-sidebar-control"])

            # G — Sampling strategy
            with gr.Row():
                sample_strategy = gr.Dropdown(
                    choices=_SAMPLE_STRATEGIES, value="TOP N",
                    label=t("dc_sample_strategy"),
                    elem_classes=["st-sidebar-control"])
                stratified_col_input = gr.Textbox(
                    label=t("dc_stratified_col"),
                    placeholder=t("dc_stratified_col_hint"),
                    visible=False,
                    elem_classes=["st-sidebar-control"])

            # Feature G — Data masking
            masking_checkbox = gr.Checkbox(label=t("dc_masking_enabled"), value=True)

            # Feature A — Incremental comparison
            with gr.Accordion(t("dc_incremental"), open=False) as incremental_accordion:
                watermark_col_input = gr.Textbox(
                    label=t("dc_watermark_col"),
                    placeholder=t("dc_watermark_hint"),
                    elem_classes=["st-sidebar-control"])
                watermark_val_input = gr.Textbox(
                    label=t("dc_watermark_value"),
                    placeholder="2024-01-01 00:00:00",
                    elem_classes=["st-sidebar-control"])
                incremental_btn = gr.Button(
                    t("dc_incremental_compare"), size="sm",
                    elem_classes=["st-connect-btn"])

            # Feature B — Quality rules
            with gr.Accordion(t("dc_quality"), open=False) as quality_accordion:
                quality_rules_input = gr.Textbox(
                    label=t("dc_quality_rules"),
                    placeholder=t("dc_quality_rules_hint"),
                    elem_classes=["st-sidebar-control"])
                quality_btn = gr.Button(
                    t("dc_quality_check"), size="sm",
                    elem_classes=["st-connect-btn"])

            # BB — Data Skew
            with gr.Accordion(t("dc_skew"), open=False) as skew_accordion:
                skew_cols_input = gr.Textbox(
                    label=t("dc_skew_column"),
                    placeholder=t("dc_skew_columns_hint"),
                    elem_classes=["st-sidebar-control"])
                skew_btn = gr.Button(
                    t("dc_skew_analyze"), size="sm",
                    elem_classes=["st-connect-btn"])

            with gr.Accordion(t("dc_checksum"), open=False) as checksum_accordion:
                checksum_cols_input = gr.Textbox(
                    label=t("dc_checksum"),
                    placeholder=t("dc_checksum_columns_hint"),
                    elem_classes=["st-sidebar-control"])
                checksum_btn = gr.Button(
                    t("dc_checksum_analyze"), size="sm",
                    elem_classes=["st-connect-btn"])

            with gr.Accordion(t("dc_partition"), open=False) as partition_accordion:
                partition_col_input = gr.Textbox(
                    label=t("dc_partition_col"),
                    placeholder=t("dc_partition_col_hint"),
                    elem_classes=["st-sidebar-control"])
                partition_btn = gr.Button(
                    t("dc_partition_analyze"), size="sm",
                    elem_classes=["st-connect-btn"])

            with gr.Accordion(t("dc_custom_agg"), open=False) as custom_agg_accordion:
                custom_agg_input = gr.Textbox(
                    label=t("dc_custom_agg"),
                    placeholder=t("dc_custom_agg_hint"),
                    lines=3,
                    elem_classes=["st-sidebar-control"])
                custom_agg_btn = gr.Button(
                    t("dc_custom_agg_analyze"), size="sm",
                    elem_classes=["st-connect-btn"])

            gr.Markdown("---")
            with gr.Row():
                schema_btn = gr.Button(t("dc_compare_schema"), size="sm", elem_classes=["st-connect-btn"])
                count_btn = gr.Button(t("dc_compare_count"), size="sm", elem_classes=["st-connect-btn"])
                sample_btn = gr.Button(t("dc_compare_sample"), size="sm", elem_classes=["st-connect-btn"])
                agg_btn = gr.Button(t("dc_compare_agg"), size="sm", elem_classes=["st-connect-btn"])
            with gr.Row():
                all_btn = gr.Button(t("dc_compare_all"), variant="primary", size="sm",
                                    elem_classes=["st-connect-btn"])
                batch_btn = gr.Button(t("dc_batch_count"), size="sm",
                                      elem_classes=["st-connect-btn"])
                profile_btn = gr.Button(t("dc_profile"), size="sm",
                                        elem_classes=["st-connect-btn"])
                batch_full_btn = gr.Button(t("dc_batch_full"), size="sm",
                                            elem_classes=["st-connect-btn"])

            # Custom SQL
            with gr.Accordion(t("dc_custom_sql"), open=False) as sql_accordion:
                sql_a_input = gr.Textbox(label="SQL (A)", placeholder=t("dc_sql_placeholder"),
                                          lines=3, elem_classes=["st-sidebar-control"])
                sql_b_input = gr.Textbox(label="SQL (B)", placeholder=t("dc_sql_placeholder"),
                                          lines=3, elem_classes=["st-sidebar-control"])
                sql_btn = gr.Button(t("dc_compare_sql"), size="sm", elem_classes=["st-connect-btn"])

            # C — Sync config + Diff SQL
            with gr.Accordion(t("dc_gen_sync"), open=False) as sync_accordion:
                gen_sync_btn = gr.Button(t("dc_gen_sync"), size="sm",
                                          elem_classes=["st-connect-btn"])
                gen_diff_sql_btn = gr.Button(t("dc_gen_diff_sql"), size="sm",
                                              elem_classes=["st-connect-btn"])

            # D — Templates
            with gr.Accordion(t("dc_templates"), open=False) as templates_accordion:
                template_name_input = gr.Textbox(
                    label=t("dc_template_name"),
                    elem_classes=["st-sidebar-control"])
                template_save_btn = gr.Button(
                    t("dc_template_save"), size="sm",
                    elem_classes=["st-connect-btn"])
                template_dd = gr.Dropdown(
                    choices=[], label=t("dc_template_load"),
                    elem_classes=["st-sidebar-control"])
                with gr.Row():
                    template_load_btn = gr.Button(
                        t("dc_template_load"), size="sm",
                        elem_classes=["st-connect-btn"])
                    template_delete_btn = gr.Button(
                        "🗑", size="sm",
                        elem_classes=["st-connect-btn"])
                template_status = gr.HTML("")
                # F6 — Batch template orchestration
                batch_template_dd = gr.Dropdown(
                    choices=[], label=t("dc_batch_templates"),
                    multiselect=True,
                    elem_classes=["st-sidebar-control"])
                batch_template_btn = gr.Button(
                    t("dc_batch_templates_run"), size="sm",
                    elem_classes=["st-connect-btn"])

            # E — Schedule
            with gr.Accordion(t("dc_schedule"), open=False) as schedule_accordion:
                schedule_interval = gr.Dropdown(
                    choices=["1", "5", "15", "30", "60"], value="15",
                    label=t("dc_schedule_interval"),
                    elem_classes=["st-sidebar-control"])
                with gr.Row():
                    schedule_start_btn = gr.Button(t("dc_schedule_start"), size="sm",
                                                    elem_classes=["st-connect-btn"])
                    schedule_stop_btn = gr.Button(t("dc_schedule_stop"), size="sm",
                                                   elem_classes=["st-connect-btn"])
                schedule_status = gr.HTML("")

            # X — Webhook
            with gr.Accordion(t("dc_webhook"), open=False) as webhook_accordion:
                webhook_url_input = gr.Textbox(
                    label=t("dc_webhook_url"),
                    placeholder=t("dc_webhook_url_hint"),
                    elem_classes=["st-sidebar-control"])
                webhook_on_fail = gr.Checkbox(
                    label=t("dc_webhook_on_fail"), value=False)

            # H — Trends
            with gr.Accordion(t("dc_trend"), open=False) as trend_accordion:
                alert_rules_input = gr.Textbox(
                    label=t("dc_alert_rules"),
                    placeholder=t("dc_alert_rules_hint"),
                    elem_classes=["st-sidebar-control"])
                trend_btn = gr.Button(t("dc_trend_chart"), size="sm",
                                       elem_classes=["st-connect-btn"])
                # Report diff (Phase 5A)
                with gr.Row():
                    report_old_dd = gr.Dropdown(choices=[], label=t("dc_report_old"),
                                                 elem_classes=["st-sidebar-control"])
                    report_new_dd = gr.Dropdown(choices=[], label=t("dc_report_new"),
                                                 elem_classes=["st-sidebar-control"])
                report_diff_btn = gr.Button(t("dc_report_compare"), size="sm",
                                             elem_classes=["st-connect-btn"])

            # Data lineage (Phase 5C)
            with gr.Accordion(t("dc_lineage_title"), open=False) as lineage_accordion:
                lineage_sql_input = gr.Textbox(
                    label=t("dc_lineage_hint"),
                    placeholder="SELECT * FROM source_table ...",
                    lines=3,
                    elem_classes=["st-sidebar-control"])
                lineage_btn = gr.Button(t("dc_lineage_title"), size="sm",
                                         elem_classes=["st-connect-btn"])

            # AA — Multi-environment
            with gr.Accordion(t("dc_environment"), open=False) as env_accordion:
                with gr.Row():
                    env_input_a = gr.Textbox(
                        label=t("dc_env_a"),
                        placeholder=t("dc_env_hint"),
                        elem_classes=["st-sidebar-control"])
                    env_input_b = gr.Textbox(
                        label=t("dc_env_b"),
                        placeholder=t("dc_env_hint"),
                        elem_classes=["st-sidebar-control"])
                env_table_input = gr.Textbox(
                    label=t("dc_select_table"),
                    placeholder="orders",
                    elem_classes=["st-sidebar-control"])
                env_compare_btn = gr.Button(
                    t("dc_compare_envs"), size="sm",
                    elem_classes=["st-connect-btn"])

            with gr.Row():
                csv_btn = gr.DownloadButton(t("dc_export_csv"), size="sm",
                                             elem_classes=["st-connect-btn"])
                excel_btn = gr.DownloadButton(t("dc_export_excel"), size="sm",
                                               elem_classes=["st-connect-btn"])
                view_btn = gr.Button(t("dc_view_report"), size="sm", variant="secondary",
                                     elem_classes=["st-connect-btn"])

            # Report persistence
            with gr.Row():
                save_btn = gr.Button(t("dc_save_report"), size="sm", elem_classes=["st-connect-btn"])
                load_dd = gr.Dropdown(choices=[], label=t("dc_load_report"),
                                       elem_classes=["st-sidebar-control"])
                load_btn = gr.Button("↻", size="sm", elem_classes=["st-connect-btn"])

        # ── Right main area ──
        with gr.Column(scale=1, elem_classes=["st-main"]):
            with gr.Row(elem_classes=["st-topbar-row"]):
                gr.HTML('<div class="st-topbar-spacer"></div>')
                home_btn = gr.Button("\U0001f3e0", size="sm", elem_classes=["st-home-btn"])
                lang_dd = gr.Dropdown(
                    choices=["English", "中文"], value="English",
                    show_label=False, container=False, min_width=140,
                    elem_classes=["st-lang-dd"],
                )
            save_status = gr.HTML(value="", visible=True)
            result_html = gr.HTML(
                value=(
                    '<div style="display:flex;align-items:center;justify-content:center;'
                    f'height:60vh;color:#9ca3af;font-size:15px;">'
                    f'{t("dc_placeholder")}</div>'
                ),
                elem_classes=["st-schema-card"],
            )
            hidden_report = gr.Textbox(visible=False, elem_id="dc-hidden-report")
            trend_plot = gr.Plot(label="", visible=False)

    # ── Wire callbacks ──

    # DS type changes
    ds_a.change(fn=_on_ds_change, inputs=[ds_a],
                outputs=[port_a, db_a, host_a, user_a, pwd_a])
    ds_b.change(fn=_on_ds_change, inputs=[ds_b],
                outputs=[port_b, db_b, host_b, user_b, pwd_b])

    # Connect buttons
    conn_a.click(
        fn=lambda *args: _do_connect(*args, side="a"),
        inputs=[ds_a, host_a, port_a, db_a, user_a, pwd_a, lang_state],
        outputs=[status_a, table_a, host_a, port_a, db_a],
    )
    conn_b.click(
        fn=lambda *args: _do_connect(*args, side="b"),
        inputs=[ds_b, host_b, port_b, db_b, user_b, pwd_b, lang_state],
        outputs=[status_b, table_b, host_b, port_b, db_b],
    )

    # G — Table search filters
    search_a.change(fn=lambda s: _filter_tables(s, "a"), inputs=[search_a], outputs=[table_a])
    search_b.change(fn=lambda s: _filter_tables(s, "b"), inputs=[search_b], outputs=[table_b])

    # Strategy change — show/hide stratified column input
    sample_strategy.change(
        fn=lambda s: gr.update(visible=(s == "STRATIFIED")),
        inputs=[sample_strategy],
        outputs=[stratified_col_input])

    # Compare buttons — pass threshold, strategy, mapping where applicable
    schema_btn.click(fn=_compare_schema,
                     inputs=[table_a, table_b, lang_state, where_input],
                     outputs=[result_html])
    count_btn.click(fn=_compare_count,
                    inputs=[table_a, table_b, lang_state, where_input, threshold_input],
                    outputs=[result_html])
    sample_btn.click(fn=_compare_sample,
                     inputs=[table_a, table_b, lang_state, where_input, key_input,
                             sample_strategy, mapping_input, masking_checkbox,
                             stratified_col_input],
                     outputs=[result_html])
    agg_btn.click(fn=_compare_agg,
                  inputs=[table_a, table_b, lang_state, where_input],
                  outputs=[result_html])
    all_btn.click(fn=_compare_all,
                  inputs=[table_a, table_b, lang_state, where_input, key_input,
                          threshold_input, sample_strategy, mapping_input,
                          masking_checkbox, webhook_url_input, webhook_on_fail,
                          skew_cols_input, stratified_col_input,
                          checksum_cols_input, partition_col_input, custom_agg_input],
                  outputs=[result_html])
    batch_btn.click(fn=_batch_count,
                    inputs=[lang_state, where_input],
                    outputs=[result_html])
    profile_btn.click(fn=_compare_profile,
                      inputs=[table_a, table_b, lang_state, where_input],
                      outputs=[result_html])
    skew_btn.click(fn=_compare_skew,
                   inputs=[table_a, table_b, lang_state, where_input, skew_cols_input],
                   outputs=[result_html])
    checksum_btn.click(fn=_compare_checksum,
                       inputs=[table_a, table_b, lang_state, where_input, checksum_cols_input],
                       outputs=[result_html])
    partition_btn.click(fn=_compare_partition,
                        inputs=[table_a, table_b, lang_state, where_input, partition_col_input],
                        outputs=[result_html])
    custom_agg_btn.click(fn=_compare_custom_agg,
                         inputs=[table_a, table_b, lang_state, where_input, custom_agg_input],
                         outputs=[result_html])
    batch_full_btn.click(fn=_batch_full,
                         inputs=[lang_state, where_input, threshold_input],
                         outputs=[result_html])

    # B — Custom SQL
    sql_btn.click(fn=_compare_sql,
                  inputs=[sql_a_input, sql_b_input, lang_state],
                  outputs=[result_html])

    # C — Sync config + Diff SQL
    gen_sync_btn.click(fn=_gen_sync_config,
                       inputs=[table_a, table_b, mapping_input, lang_state],
                       outputs=[result_html])
    gen_diff_sql_btn.click(fn=_gen_diff_sql_fn,
                           inputs=[table_a, table_b, lang_state],
                           outputs=[result_html])

    # Feature A — Incremental
    incremental_btn.click(
        fn=_incremental_fn,
        inputs=[table_a, table_b, lang_state, where_input,
                watermark_col_input, watermark_val_input, key_input,
                mapping_input, masking_checkbox],
        outputs=[result_html],
    )

    # Feature B — Quality
    quality_btn.click(
        fn=_check_quality_fn,
        inputs=[table_a, table_b, lang_state, where_input, quality_rules_input],
        outputs=[result_html],
    )

    # Feature D — Templates
    template_save_btn.click(
        fn=_save_template,
        inputs=[template_name_input, table_a, table_b, where_input, key_input,
                threshold_input, mapping_input, sample_strategy, watermark_col_input,
                quality_rules_input, webhook_url_input, webhook_on_fail, masking_checkbox,
                lang_state],
        outputs=[template_status, template_dd],
    )
    template_load_btn.click(
        fn=_load_template,
        inputs=[template_dd, lang_state],
        outputs=[table_a, table_b, where_input, key_input, threshold_input,
                 mapping_input, sample_strategy, watermark_col_input,
                 quality_rules_input, webhook_url_input,
                 webhook_on_fail, masking_checkbox],
    )
    template_delete_btn.click(
        fn=_delete_template,
        inputs=[template_dd, lang_state],
        outputs=[template_status, template_dd],
    )

    # F6 — Batch template orchestration
    def _refresh_batch_template_choices():
        names = [d["name"] for d in _TEMPLATES_STORE.list()]
        return gr.update(choices=names)

    batch_template_dd.focus(fn=_refresh_batch_template_choices, outputs=[batch_template_dd])
    batch_template_btn.click(
        fn=_batch_templates_fn,
        inputs=[batch_template_dd, lang_state],
        outputs=[result_html],
    )

    # Feature H — Environment compare
    env_compare_btn.click(
        fn=_compare_environments_fn,
        inputs=[env_input_a, env_input_b, env_table_input, lang_state, where_input],
        outputs=[result_html],
    )

    # E — Schedule
    schedule_start_btn.click(
        fn=_schedule_start,
        inputs=[table_a, table_b, lang_state, where_input, key_input, schedule_interval],
        outputs=[schedule_status],
    )
    schedule_stop_btn.click(fn=_schedule_stop, inputs=[lang_state],
                            outputs=[schedule_status])

    # H — Trends
    trend_btn.click(fn=_show_trend, inputs=[lang_state, alert_rules_input],
                    outputs=[result_html, trend_plot])

    # Report diff (Phase 5A)
    def _refresh_report_dds():
        if not _REPORTS_DIR.is_dir():
            return gr.update(choices=[]), gr.update(choices=[])
        files = sorted(_REPORTS_DIR.glob("compare_*.json"), reverse=True)
        names = [f.name for f in files[:_MAX_REPORT_FILES]]
        return gr.update(choices=names), gr.update(choices=names)

    report_diff_btn.click(fn=_compare_reports_fn,
                          inputs=[report_old_dd, report_new_dd, lang_state],
                          outputs=[result_html])
    report_old_dd.focus(fn=_refresh_report_dds,
                        outputs=[report_old_dd, report_new_dd])

    # Data lineage (Phase 5C)
    lineage_btn.click(fn=_show_lineage_fn,
                      inputs=[table_a, lang_state, lineage_sql_input],
                      outputs=[result_html])

    # A — Presets
    preset_save_btn.click(
        fn=_save_preset,
        inputs=[preset_name_input, ds_a, host_a, port_a, db_a, user_a, pwd_a, preset_env_input, lang_state],
        outputs=[preset_status],
    )
    preset_save_b_btn.click(
        fn=_save_preset,
        inputs=[preset_name_input, ds_b, host_b, port_b, db_b, user_b, pwd_b, preset_env_input, lang_state],
        outputs=[preset_status],
    )
    preset_dd_a.change(fn=_load_preset, inputs=[preset_dd_a, lang_state],
                       outputs=[ds_a, host_a, port_a, db_a, user_a, pwd_a, preset_env_input])
    preset_dd_b.change(fn=_load_preset, inputs=[preset_dd_b, lang_state],
                       outputs=[ds_b, host_b, port_b, db_b, user_b, pwd_b, preset_env_input])

    # Export buttons
    csv_btn.click(fn=_export_csv, inputs=[lang_state], outputs=[csv_btn])
    excel_btn.click(fn=_export_excel, inputs=[lang_state], outputs=[excel_btn])

    # View Report — open in new tab
    def _build_report_page(lang_val):
        with holder_lock:
            report = holder.get("last_report")
        if not report:
            return dc(lang_val, "dc_no_report")
        return build_standalone_report(report, lang_val)

    view_btn.click(
        fn=_build_report_page,
        inputs=[lang_state],
        outputs=[hidden_report],
    ).then(
        fn=None,
        inputs=[hidden_report],
        js="""(html) => {
            if (!html || html.length < 50) { alert(html); return; }
            const w = window.open('', '_blank');
            if (w) { w.document.write(html); w.document.close(); }
        }""",
    )

    # E — Report persistence
    save_btn.click(fn=_save_report, inputs=[lang_state], outputs=[save_status])
    load_btn.click(fn=_list_reports, inputs=[], outputs=[load_dd])
    load_dd.change(fn=_load_report, inputs=[load_dd, lang_state], outputs=[result_html])

    # Home button
    home_btn.click(fn=None, js="() => { window.location.href = '/'; }")

    # Language switch — update all component labels/text
    def _switch_lang(choice):
        lg = "zh" if choice == "中文" else "en"
        t_fn = lambda k: dc(lg, k)
        return (
            lg,                                                         # lang_state
            t_fn("dc_title"),                                           # title_md
            f"**{t_fn('dc_source_a')}**",                               # src_a_md
            f"**{t_fn('dc_source_b')}**",                               # src_b_md
            gr.update(label=t_fn("dc_datasource_type")),                # ds_a
            gr.update(label=t_fn("dc_datasource_type")),                # ds_b
            gr.update(label=t_fn("dc_host")),                           # host_a
            gr.update(label=t_fn("dc_host")),                           # host_b
            gr.update(label=t_fn("dc_port")),                           # port_a
            gr.update(label=t_fn("dc_port")),                           # port_b
            gr.update(label=t_fn("dc_database")),                       # db_a
            gr.update(label=t_fn("dc_database")),                       # db_b
            gr.update(label=t_fn("dc_username")),                       # user_a
            gr.update(label=t_fn("dc_username")),                       # user_b
            gr.update(label=t_fn("dc_password")),                       # pwd_a
            gr.update(label=t_fn("dc_password")),                       # pwd_b
            gr.update(value=t_fn("dc_connect")),                        # conn_a
            gr.update(value=t_fn("dc_connect")),                        # conn_b
            gr.update(label=t_fn("dc_status")),                         # status_a
            gr.update(label=t_fn("dc_status")),                         # status_b
            gr.update(placeholder=t_fn("dc_search_tables")),            # search_a
            gr.update(placeholder=t_fn("dc_search_tables")),            # search_b
            gr.update(label=t_fn("dc_select_table")),                   # table_a
            gr.update(label=t_fn("dc_select_table")),                   # table_b
            gr.update(label=t_fn("dc_where_clause"),
                      placeholder=t_fn("dc_where_hint")),               # where_input
            gr.update(label=t_fn("dc_key_columns")),                    # key_input
            gr.update(value=t_fn("dc_compare_schema")),                 # schema_btn
            gr.update(value=t_fn("dc_compare_count")),                  # count_btn
            gr.update(value=t_fn("dc_compare_sample")),                 # sample_btn
            gr.update(value=t_fn("dc_compare_agg")),                    # agg_btn
            gr.update(value=t_fn("dc_compare_all")),                    # all_btn
            gr.update(value=t_fn("dc_batch_count")),                    # batch_btn
            gr.update(value=t_fn("dc_profile")),                        # profile_btn
            gr.update(label=t_fn("dc_custom_sql")),                     # sql_accordion
            gr.update(placeholder=t_fn("dc_sql_placeholder")),          # sql_a_input
            gr.update(placeholder=t_fn("dc_sql_placeholder")),          # sql_b_input
            gr.update(value=t_fn("dc_compare_sql")),                    # sql_btn
            gr.update(label=t_fn("dc_export_csv")),                     # csv_btn
            gr.update(label=t_fn("dc_export_excel")),                   # excel_btn
            gr.update(value=t_fn("dc_save_report")),                    # save_btn
            gr.update(label=t_fn("dc_load_report")),                    # load_dd
            gr.update(value=t_fn("dc_view_report")),                    # view_btn
            # Round 4 components
            gr.update(label=t_fn("dc_presets")),                        # presets_accordion
            gr.update(label=t_fn("dc_preset_name")),                    # preset_name_input
            gr.update(label=t_fn("dc_preset_env")),                     # preset_env_input
            gr.update(value=f'{t_fn("dc_preset_save")} (A)'),           # preset_save_btn
            gr.update(value=f'{t_fn("dc_preset_save")} (B)'),           # preset_save_b_btn
            gr.update(label=f"{t_fn('dc_preset_load')} (A)"),           # preset_dd_a
            gr.update(label=f"{t_fn('dc_preset_load')} (B)"),           # preset_dd_b
            gr.update(label=t_fn("dc_threshold"),
                      placeholder=t_fn("dc_threshold_hint")),           # threshold_input
            gr.update(label=t_fn("dc_col_mapping"),
                      placeholder=t_fn("dc_col_mapping_hint")),         # mapping_input
            gr.update(label=t_fn("dc_sample_strategy")),                # sample_strategy
            gr.update(value=t_fn("dc_batch_full")),                     # batch_full_btn
            gr.update(label=t_fn("dc_gen_sync")),                       # sync_accordion
            gr.update(value=t_fn("dc_gen_sync")),                       # gen_sync_btn
            gr.update(label=t_fn("dc_schedule")),                       # schedule_accordion
            gr.update(label=t_fn("dc_schedule_interval")),              # schedule_interval
            gr.update(value=t_fn("dc_schedule_start")),                 # schedule_start_btn
            gr.update(value=t_fn("dc_schedule_stop")),                  # schedule_stop_btn
            gr.update(label=t_fn("dc_trend")),                          # trend_accordion
            gr.update(label=t_fn("dc_alert_rules"),
                      placeholder=t_fn("dc_alert_rules_hint")),        # alert_rules_input
            gr.update(value=t_fn("dc_trend_chart")),                    # trend_btn
            # Round 5 components
            gr.update(label=t_fn("dc_masking_enabled")),                # masking_checkbox
            gr.update(label=t_fn("dc_incremental")),                    # incremental_accordion
            gr.update(label=t_fn("dc_watermark_col"),
                      placeholder=t_fn("dc_watermark_hint")),           # watermark_col_input
            gr.update(label=t_fn("dc_watermark_value")),                # watermark_val_input
            gr.update(value=t_fn("dc_incremental_compare")),            # incremental_btn
            gr.update(label=t_fn("dc_quality")),                        # quality_accordion
            gr.update(label=t_fn("dc_quality_rules"),
                      placeholder=t_fn("dc_quality_rules_hint")),       # quality_rules_input
            gr.update(value=t_fn("dc_quality_check")),                  # quality_btn
            gr.update(value=t_fn("dc_gen_diff_sql")),                   # gen_diff_sql_btn
            gr.update(label=t_fn("dc_templates")),                      # templates_accordion
            gr.update(label=t_fn("dc_template_name")),                  # template_name_input
            gr.update(value=t_fn("dc_template_save")),                  # template_save_btn
            gr.update(label=t_fn("dc_template_load")),                  # template_dd
            gr.update(value=t_fn("dc_template_load")),                  # template_load_btn
            gr.update(label=t_fn("dc_webhook")),                        # webhook_accordion
            gr.update(label=t_fn("dc_webhook_url"),
                      placeholder=t_fn("dc_webhook_url_hint")),         # webhook_url_input
            gr.update(label=t_fn("dc_webhook_on_fail")),                # webhook_on_fail
            gr.update(label=t_fn("dc_environment")),                    # env_accordion
            gr.update(label=t_fn("dc_env_a"),
                      placeholder=t_fn("dc_env_hint")),                 # env_input_a
            gr.update(label=t_fn("dc_env_b"),
                      placeholder=t_fn("dc_env_hint")),                 # env_input_b
            gr.update(label=t_fn("dc_select_table")),                   # env_table_input
            gr.update(value=t_fn("dc_compare_envs")),                   # env_compare_btn
            # Round 6 — skew
            gr.update(label=t_fn("dc_skew")),                           # skew_accordion
            gr.update(label=t_fn("dc_skew_column"),
                      placeholder=t_fn("dc_skew_columns_hint")),        # skew_cols_input
            gr.update(value=t_fn("dc_skew_analyze")),                   # skew_btn
            # Round 7 — checksum, partition, custom agg
            gr.update(label=t_fn("dc_checksum")),                       # checksum_accordion
            gr.update(label=t_fn("dc_checksum"),
                      placeholder=t_fn("dc_checksum_columns_hint")),    # checksum_cols_input
            gr.update(value=t_fn("dc_checksum_analyze")),               # checksum_btn
            gr.update(label=t_fn("dc_partition")),                      # partition_accordion
            gr.update(label=t_fn("dc_partition_col"),
                      placeholder=t_fn("dc_partition_col_hint")),       # partition_col_input
            gr.update(value=t_fn("dc_partition_analyze")),              # partition_btn
            gr.update(label=t_fn("dc_custom_agg")),                    # custom_agg_accordion
            gr.update(label=t_fn("dc_custom_agg"),
                      placeholder=t_fn("dc_custom_agg_hint")),         # custom_agg_input
            gr.update(value=t_fn("dc_custom_agg_analyze")),            # custom_agg_btn
            # Round 8 — stratified sampling
            gr.update(label=t_fn("dc_stratified_col"),
                      placeholder=t_fn("dc_stratified_col_hint")),     # stratified_col_input
            # Round 9 — batch templates
            gr.update(label=t_fn("dc_batch_templates")),               # batch_template_dd
            gr.update(value=t_fn("dc_batch_templates_run")),           # batch_template_btn
            # Round 10 — report diff + lineage
            gr.update(label=t_fn("dc_report_old")),                    # report_old_dd
            gr.update(label=t_fn("dc_report_new")),                    # report_new_dd
            gr.update(value=t_fn("dc_report_compare")),                # report_diff_btn
            gr.update(label=t_fn("dc_lineage_title")),                 # lineage_accordion
            gr.update(label=t_fn("dc_lineage_hint")),                  # lineage_sql_input
            gr.update(value=t_fn("dc_lineage_title")),                 # lineage_btn
        )

    lang_dd.change(
        fn=_switch_lang,
        inputs=[lang_dd],
        outputs=[
            lang_state, title_md, src_a_md, src_b_md,
            ds_a, ds_b, host_a, host_b, port_a, port_b, db_a, db_b,
            user_a, user_b, pwd_a, pwd_b,
            conn_a, conn_b, status_a, status_b,
            search_a, search_b, table_a, table_b,
            where_input, key_input,
            schema_btn, count_btn, sample_btn, agg_btn,
            all_btn, batch_btn, profile_btn,
            sql_accordion, sql_a_input, sql_b_input, sql_btn,
            csv_btn, excel_btn, save_btn, load_dd, view_btn,
            # Round 4
            presets_accordion, preset_name_input, preset_env_input,
            preset_save_btn, preset_save_b_btn,
            preset_dd_a, preset_dd_b,
            threshold_input, mapping_input, sample_strategy,
            batch_full_btn,
            sync_accordion, gen_sync_btn,
            schedule_accordion, schedule_interval,
            schedule_start_btn, schedule_stop_btn,
            trend_accordion, alert_rules_input, trend_btn,
            # Round 5
            masking_checkbox,
            incremental_accordion, watermark_col_input, watermark_val_input, incremental_btn,
            quality_accordion, quality_rules_input, quality_btn,
            gen_diff_sql_btn,
            templates_accordion, template_name_input, template_save_btn,
            template_dd, template_load_btn,
            webhook_accordion, webhook_url_input, webhook_on_fail,
            env_accordion, env_input_a, env_input_b, env_table_input, env_compare_btn,
            # Round 6 — skew
            skew_accordion, skew_cols_input, skew_btn,
            # Round 7 — checksum, partition, custom agg
            checksum_accordion, checksum_cols_input, checksum_btn,
            partition_accordion, partition_col_input, partition_btn,
            custom_agg_accordion, custom_agg_input, custom_agg_btn,
            # Round 8 — stratified sampling
            stratified_col_input,
            # Round 9 — batch templates
            batch_template_dd, batch_template_btn,
            # Round 10 — report diff + lineage
            report_old_dd, report_new_dd, report_diff_btn,
            lineage_accordion, lineage_sql_input, lineage_btn,
        ],
    )
